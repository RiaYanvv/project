from pathlib import Path
from openpyxl import load_workbook
import argparse


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-root", type=Path, required=True)
    parser.add_argument("--excel", required=True)
    args = parser.parse_args()

    data_root = args.data_root.expanduser().resolve()
    excel_path = data_root / args.excel

    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    sheet = workbook["evidence"] if "evidence" in workbook.sheetnames else workbook.active

    rows = sheet.iter_rows(values_only=True)
    headers = [clean(x) for x in next(rows)]

    raw_root = data_root / "raw"

    # Build index: evidence_id -> actual files found anywhere under raw/
    file_index = {}

    for path in raw_root.rglob("*"):
        if not path.is_file():
            continue

        if path.name == ".gitkeep":
            continue

        # Correctly handle double extensions such as FR_001.json.gz
        if path.name.lower().endswith(".json.gz"):
            evidence_id = path.name[:-8]
        else:
            evidence_id = path.stem

        file_index.setdefault(evidence_id, []).append(path)
    total = 0
    correct = 0
    mismatch = 0
    missing = 0

    for values in rows:
        record = {
            headers[i]: clean(values[i])
            for i in range(min(len(headers), len(values)))
            if headers[i]
        }

        evidence_id = record.get("evidence_id", "")
        document_path = record.get("document_path", "")

        if not evidence_id:
            continue

        total += 1

        expected = None

        if document_path:
            path = Path(document_path)

            if path.parts and path.parts[0] == "data":
                path = Path(*path.parts[1:])

            expected = data_root / path

        if expected and expected.exists():
            print(f"[OK] {evidence_id}")
            correct += 1
            continue

        candidates = file_index.get(evidence_id, [])

        if candidates:
            mismatch += 1

            print(f"\n[MISMATCH] {evidence_id}")
            print(f"  Excel:  {document_path or '(empty)'}")

            for candidate in candidates:
                relative = candidate.relative_to(data_root)
                print(f"  Actual: data/{relative}")

        else:
            missing += 1

            print(f"\n[MISSING] {evidence_id}")
            print(f"  Excel: {document_path or '(empty)'}")
            print("  Actual: no matching file found under data/raw/")

    print()
    print("=== PATH CHECK COMPLETE ===")
    print(f"Total:    {total}")
    print(f"Correct:  {correct}")
    print(f"Mismatch: {mismatch}")
    print(f"Missing:  {missing}")


if __name__ == "__main__":
    main()
