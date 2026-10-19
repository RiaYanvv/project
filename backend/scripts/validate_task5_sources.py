"""Validate Task 5 case-source traceability and URL reachability.

The case database is an evaluation answer key, so its external sources must be
independently auditable. This script checks both the workbook wiring and the
live URLs, then writes:

    data/evaluation/task5_source_validation.xlsx

Exit code 0 means: no structural error and no confirmed dead link. Some Chinese
news sites sit behind WAF/captcha or render the article with JavaScript; those
rows are reported explicitly as OK_WAF / OK_RENDER instead of being silently
labelled as fully fetched.
"""
from __future__ import annotations

import argparse
import concurrent.futures as cf
import os
import re
import subprocess
import tempfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_CASE_DB = Path("/Users/zhengyuanyuan/Desktop/project-data/data/processed/case_database_v0.1.xlsx")
DEFAULT_OUT = Path("/Users/zhengyuanyuan/Desktop/project-data/data/evaluation/task5_source_validation.xlsx")
UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 Chrome/124.0 Safari/537.36"
ERROR_TITLES = ("页面没有找到", "页面不存在", "404 not found", "内容不存在", "文章不存在", "出错啦")
WAF_MARKERS = ("resverify", "tcaptcha", "captcha", "__tst_status", "访问验证", "人机验证")


def norm(value) -> str:
    return str(value or "").strip()


def read_rows(path: Path, sheet: str) -> list[dict]:
    ws = load_workbook(path, read_only=True, data_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [norm(c) for c in rows[0]]
    return [dict(zip(hdr, row)) for row in rows[1:] if row and row[0]]


def extract_title(body: bytes) -> str:
    for pattern in (
        rb"<title[^>]*>(.*?)</title>",
        rb"<meta[^>]+property=[\"']og:title[\"'][^>]+content=[\"'](.*?)[\"']",
        rb"<meta[^>]+content=[\"'](.*?)[\"'][^>]+property=[\"']og:title[\"']",
    ):
        match = re.search(pattern, body[:400000], re.I | re.S)
        if not match:
            continue
        raw = match.group(1)
        for enc in ("utf-8", "gb18030"):
            try:
                text = raw.decode(enc)
            except UnicodeDecodeError:
                continue
            return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", text)).strip()[:220]
    return ""


def check_url(row: dict) -> dict:
    url = norm(row.get("url"))
    fd, tmp = tempfile.mkstemp()
    os.close(fd)
    try:
        cp = subprocess.run(
            [
                "curl", "-L", "-sS", "--max-time", "25", "--connect-timeout", "10",
                "-A", UA, "-H", "Accept-Language: zh-CN,zh;q=0.9,en;q=0.8",
                "-w", "%{http_code}\t%{url_effective}\t%{size_download}",
                "-o", tmp, url,
            ],
            capture_output=True,
            text=True,
            timeout=35,
        )
        meta = cp.stdout.strip().split("\t")
        code = meta[0] if meta and meta[0] else "ERR"
        final_url = meta[1] if len(meta) > 1 else ""
        size = int(meta[2]) if len(meta) > 2 and meta[2].isdigit() else 0
        body = Path(tmp).read_bytes()[:500000]
        page_title = extract_title(body)
        raw = body.decode("utf-8", "ignore")
        lower = raw.lower()
        dead = (
            not code.startswith("2")
            or any(x.lower() in page_title.lower() for x in ERROR_TITLES)
            or any(x.lower() in lower for x in ("404 not found", "page not found"))
        )
        waf = any(x.lower() in lower for x in WAF_MARKERS)
        if dead:
            verdict = "DEAD"
            detail = "链接失效或返回错误页"
        elif page_title:
            verdict = "OK"
            detail = "HTTP 可达且抓到页面标题"
        elif waf:
            verdict = "OK_WAF"
            detail = "HTTP 可达；站点 WAF/验证码阻断自动抓取，需浏览器人工打开"
        else:
            verdict = "OK_RENDER"
            detail = "HTTP 可达；页面由客户端渲染或站点未返回标题，需浏览器人工确认"
        return {
            "http_status": code,
            "final_url": final_url,
            "page_title": page_title,
            "content_bytes": size,
            "verdict": verdict,
            "detail": detail,
            "curl_error": cp.stderr.strip()[:300],
        }
    except Exception as exc:
        return {
            "http_status": "ERR",
            "final_url": "",
            "page_title": "",
            "content_bytes": 0,
            "verdict": "DEAD",
            "detail": f"请求失败：{type(exc).__name__}: {exc}",
            "curl_error": "",
        }
    finally:
        Path(tmp).unlink(missing_ok=True)


def style(ws, widths: list[int], header_color: str = "1F3864") -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=header_color)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--case-db", type=Path, default=DEFAULT_CASE_DB)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--offline", action="store_true", help="只做结构性校验，不访问 URL")
    args = ap.parse_args()

    cases = read_rows(args.case_db, "Case Database")
    sources = read_rows(args.case_db, "Evidence Index")
    errors: list[str] = []
    by_id = {norm(s["evidence_id"]): s for s in sources}
    by_case = {norm(c["case_id"]): c for c in cases}

    if len(by_id) != len(sources):
        errors.append("Evidence Index 存在重复 evidence_id")
    urls = [norm(s.get("url")) for s in sources]
    duplicate_urls = sorted({url for url in urls if url and urls.count(url) > 1})
    if duplicate_urls:
        errors.append(f"Evidence Index 存在重复 URL: {duplicate_urls}")

    used: dict[str, set[str]] = {sid: set() for sid in by_id}
    case_links: list[dict] = []
    for case in cases:
        cid = norm(case["case_id"])
        primary_url = norm(case.get("url"))
        if not primary_url:
            errors.append(f"{cid}: 缺少案例主来源 URL")
        refs = [x.strip() for x in re.split(r"[;,]", norm(case.get("external_source_refs"))) if x.strip()]
        if not refs:
            errors.append(f"{cid}: 缺少 external_source_refs")
        for ref in refs:
            if ref not in by_id:
                errors.append(f"{cid}: 悬空来源 {ref}")
                continue
            used[ref].add(cid)
            case_links.append({
                "case_id": cid,
                "evidence_id": ref,
                "outlet": by_id[ref].get("outlet"),
                "title": by_id[ref].get("title"),
                "date": by_id[ref].get("date"),
                "url": by_id[ref].get("url"),
            })
        resolved_urls = {norm(by_id[ref].get("url")) for ref in refs if ref in by_id}
        if primary_url and primary_url not in resolved_urls:
            errors.append(f"{cid}: 案例主来源 URL 不在 external_source_refs 的 Evidence Index 记录中")

    for sid, source in by_id.items():
        declared = {x.strip() for x in norm(source.get("used_by_cases")).split(";") if x.strip()}
        if declared != used[sid]:
            errors.append(f"{sid}: used_by_cases={sorted(declared)}，实际={sorted(used[sid])}")
        if not norm(source.get("url")):
            errors.append(f"{sid}: 缺少 URL")
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", norm(source.get("date"))):
            errors.append(f"{sid}: date 不是 YYYY-MM-DD 格式（{source.get('date')!r}）")

    results = []
    for source in sources:
        rec = dict(source)
        if args.offline:
            live = {
                "http_status": "SKIP", "final_url": "", "page_title": "",
                "content_bytes": 0, "verdict": "SKIP",
                "detail": "离线结构校验", "curl_error": "",
            }
        else:
            live = check_url(source)
        rec.update(live)
        results.append(rec)

    dead = [r for r in results if r["verdict"] == "DEAD"]
    ok = [r for r in results if r["verdict"] == "OK"]
    waf = [r for r in results if r["verdict"] == "OK_WAF"]
    render = [r for r in results if r["verdict"] == "OK_RENDER"]

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Task 5 — Source Validation", "Value", "Note"])
    rows = [
        ("validated_at", date.today().isoformat(), "审计日期"),
        ("cases", len(cases), "Task 5 案例数"),
        ("sources", len(sources), "Evidence Index 来源数"),
        ("case_source_links", len(case_links), "案例 → 来源映射数"),
        ("structural_errors", len(errors), "悬空引用 / used_by_cases 不一致 / URL 缺失或重复 / 主来源未映射 / date 格式"),
        ("http_ok_content", len(ok), "抓取到标题且未命中错误页"),
        ("http_ok_waf", len(waf), "HTTP 可达但自动抓取被 WAF/验证码阻断"),
        ("http_ok_render", len(render), "HTTP 可达但标题由客户端渲染或缺失"),
        ("dead_links", len(dead), "必须为 0"),
        ("result", "PASS" if not errors and not dead else "FAIL", "结构错误与 dead link 均须为 0"),
    ]
    for row in rows:
        summary.append(row)
    for row in summary.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
    style(summary, [26, 18, 80], "2E5E4E")

    src = wb.create_sheet("Sources")
    src.append([
        "evidence_id", "used_by_cases", "outlet", "listed_title", "date", "url",
        "http_status", "final_url", "page_title", "content_bytes", "verdict", "detail",
    ])
    for rec in results:
        src.append([
            rec.get("evidence_id"), rec.get("used_by_cases"), rec.get("outlet"),
            rec.get("title"), rec.get("date"), rec.get("url"), rec.get("http_status"),
            rec.get("final_url"), rec.get("page_title"), rec.get("content_bytes"),
            rec.get("verdict"), rec.get("detail"),
        ])
    fill = {
        "OK": PatternFill("solid", fgColor="E2EFDA"),
        "OK_WAF": PatternFill("solid", fgColor="FFF2CC"),
        "OK_RENDER": PatternFill("solid", fgColor="DDEBF7"),
        "DEAD": PatternFill("solid", fgColor="F4CCCC"),
    }
    for row in src.iter_rows(min_row=2):
        row[10].fill = fill.get(str(row[10].value), PatternFill())
    style(src, [20, 26, 34, 58, 13, 54, 12, 54, 58, 14, 14, 54])

    link_ws = wb.create_sheet("Case Links")
    link_ws.append(["case_id", "evidence_id", "outlet", "title", "date", "url"])
    for rec in case_links:
        link_ws.append([rec[k] for k in ("case_id", "evidence_id", "outlet", "title", "date", "url")])
    style(link_ws, [14, 22, 34, 60, 13, 56], "6B4C2A")

    err_ws = wb.create_sheet("Structural Errors")
    err_ws.append(["error"])
    for error in errors:
        err_ws.append([error])
    style(err_ws, [120], "9D3E3E")

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)

    print(f"cases={len(cases)} sources={len(sources)} links={len(case_links)}")
    print(f"OK={len(ok)} OK_WAF={len(waf)} OK_RENDER={len(render)} DEAD={len(dead)} errors={len(errors)}")
    if dead:
        print("dead links:")
        for rec in dead:
            print(f"  - {rec['evidence_id']}: {rec['url']} ({rec['detail']})")
    if errors:
        print("structural errors:")
        for error in errors:
            print("  -", error)
    print(f"saved {args.out}")
    return 0 if not errors and not dead else 1


if __name__ == "__main__":
    raise SystemExit(main())
