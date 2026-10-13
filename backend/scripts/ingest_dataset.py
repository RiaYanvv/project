from __future__ import annotations

import argparse
import csv
import gzip
import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from bs4 import BeautifulSoup


BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))

from app.config import Settings
from app.knowledge import KnowledgeDocument, KnowledgeRepository


def clean(value) -> str:
    """Convert Excel cell values to clean strings."""
    if value is None:
        return ""
    return str(value).strip()


def read_pdf(path: Path) -> str:
    """Extract text from a PDF."""
    reader = PdfReader(str(path))
    pages = []

    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text)

    return "\n\n".join(pages)


def read_csv(path: Path) -> str:
    """Convert CSV rows into searchable text."""
    lines = []

    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
        reader = csv.DictReader(f)

        for row in reader:
            parts = [
                f"{key}: {value}"
                for key, value in row.items()
                if key and value not in (None, "")
            ]

            if parts:
                lines.append(" | ".join(parts))

    return "\n".join(lines)


def read_xlsx(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)

    lines = []

    for sheet in workbook.worksheets:
        lines.append(f"Sheet: {sheet.title}")

        for row in sheet.iter_rows(values_only=True):
            values = [
                clean(value)
                for value in row
                if value is not None and clean(value)
            ]

            if values:
                lines.append(" | ".join(values))

    workbook.close()

    return "\n".join(lines)


def read_json(path: Path) -> str:
    """Convert JSON into searchable text."""
    with path.open("r", encoding="utf-8", errors="replace") as f:
        data = json.load(f)

    return json.dumps(data, ensure_ascii=False, indent=2)


def read_json_gz(path: Path) -> str:
    """Read a gzip JSON dataset.

    Bulk record dumps (e.g. sanctioned-party screening lists) are converted
    into a compact, de-duplicated descriptor plus a bounded name index.
    Emitting one chunk per raw record floods the retrieval index with tens of
    thousands of near-identical boilerplate passages, so we deliberately keep
    only enough structured text for the source card and name lookup.
    """
    with gzip.open(path, "rt", encoding="utf-8") as f:
        data = json.load(f)

    if isinstance(data, dict) and isinstance(data.get("results"), list):
        records = [r for r in data["results"] if isinstance(r, dict)]

        names: list[str] = []
        seen_names: set[str] = set()
        programs: dict[str, int] = {}
        sources: dict[str, int] = {}

        for record in records:
            name = record.get("name")
            if name:
                name = str(name).strip()
                if name and name not in seen_names:
                    seen_names.add(name)
                    names.append(name)

            for program in record.get("programs") or []:
                programs[str(program)] = programs.get(str(program), 0) + 1

            source = record.get("source")
            if source:
                sources[str(source)] = sources.get(str(source), 0) + 1

        top_programs = sorted(programs, key=programs.get, reverse=True)[:15]
        top_sources = sorted(sources, key=sources.get, reverse=True)[:10]

        header = [
            f"Bulk database source with {len(records)} records.",
            f"Distinct listed entity names: {len(names)}.",
        ]
        if top_programs:
            header.append("Programs covered: " + ", ".join(top_programs) + ".")
        if top_sources:
            header.append("Component lists: " + ", ".join(top_sources) + ".")

        # A bounded, de-duplicated name index keeps entity lookup usable
        # without re-introducing thousands of repetitive passages.
        index_cap = min(len(names), 600)
        name_blocks = [
            ", ".join(names[i : i + 50])
            for i in range(0, index_cap, 50)
        ]

        parts = ["\n".join(header)]
        if name_blocks:
            parts.append("Indexed entity names (sample):")
            parts.extend(name_blocks)
        return "\n\n".join(parts)

    # Fallback for other JSON structures
    return json.dumps(data, ensure_ascii=False)


def read_html(path: Path) -> str:
    """Extract article-oriented visible text from HTML."""
    import json
    import re

    raw_html = path.read_bytes()
    soup = BeautifulSoup(raw_html, "html.parser")

    # Tencent News embeds the canonical article body in
    # window.DATA["originContent"]["text"]. Prefer that structured
    # article body over rendered page chrome when available.
    for script in soup.find_all("script"):
        script_text = script.string or script.get_text()
        if "window.DATA" not in script_text:
            continue

        match = re.search(
            r"window\.DATA\s*=\s*(\{.*?\})\s*;?\s*$",
            script_text,
            flags=re.S,
        )

        if not match:
            continue

        try:
            data = json.loads(match.group(1))
        except json.JSONDecodeError:
            continue

        origin = data.get("originContent")
        if not isinstance(origin, dict):
            continue

        article_html = origin.get("text")
        if not isinstance(article_html, str) or len(article_html) < 300:
            continue

        article_soup = BeautifulSoup(article_html, "html.parser")

        for tag in article_soup([
            "script", "style", "noscript",
            "svg", "canvas",
        ]):
            tag.decompose()

        extracted = article_soup.get_text(
            separator="\n",
            strip=True,
        )

        cleaned_lines = []
        seen = set()

        for line in extracted.splitlines():
            line = " ".join(line.split()).strip()

            if not line:
                continue

            if line in seen:
                continue

            seen.add(line)
            cleaned_lines.append(line)

        return "\n".join(cleaned_lines)

    # Generic HTML fallback.
    for tag in soup([
        "script", "style", "noscript", "svg", "canvas",
        "nav", "header", "footer", "form", "button",
        "aside",
    ]):
        tag.decompose()

    noise_keywords = (
        "comment", "recommend", "related", "share",
        "sidebar", "footer", "header", "navigation",
        "advert", "toolbar",
    )

    for tag in list(soup.find_all(True)):
        if tag.parent is None:
            continue

        classes = " ".join(tag.get("class", []))
        tag_id = tag.get("id", "")
        marker = f"{classes} {tag_id}".lower()

        if any(keyword in marker for keyword in noise_keywords):
            tag.decompose()

    selectors = [
        "#article-content",
        ".article-content",
        ".content-article",
        ".article-body",
        ".article_content",
        "[class*='article-content']",
        "article",
        "main",
    ]

    content_root = None

    for selector in selectors:
        candidate = soup.select_one(selector)

        if candidate and len(candidate.get_text(strip=True)) >= 300:
            content_root = candidate
            break

    if content_root is None:
        content_root = soup.body or soup

    extracted = content_root.get_text(
        separator="\n",
        strip=True,
    )

    cleaned_lines = []
    seen = set()

    for line in extracted.splitlines():
        line = " ".join(line.split()).strip()

        if not line:
            continue

        if line in seen:
            continue

        seen.add(line)
        cleaned_lines.append(line)

    return "\n".join(cleaned_lines)

def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def extract_content(path: Path) -> str:
    # .json.gz 必须单独判断
    if path.name.lower().endswith(".json.gz"):
        return read_json_gz(path)

    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return read_pdf(path)

    if suffix == ".csv":
        return read_csv(path)
    
    if suffix == ".xlsx":
        return read_xlsx(path)

    if suffix in {".html", ".htm"}:
        return read_html(path)

    if suffix in {".txt", ".md"}:
        return read_text(path)

    raise ValueError(f"Unsupported file type: {path}")


def load_metadata(excel_path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)

    # Prefer the canonical metadata sheet. Relying on the workbook's active
    # sheet is fragile: opening the file in Excel and clicking another tab
    # silently switches ingestion to a review sheet with different columns.
    sheet = None
    for name in ("evidence", "sources"):
        if name in workbook.sheetnames:
            sheet = workbook[name]
            break
    if sheet is None:
        sheet = workbook.active

    rows = sheet.iter_rows(values_only=True)

    try:
        headers = [clean(value) for value in next(rows)]
    except StopIteration:
        return []

    records = []

    for values in rows:
        record = {
            headers[i]: clean(values[i])
            for i in range(min(len(headers), len(values)))
            if headers[i]
        }

        if record.get("evidence_id"):
            records.append(record)

    return records


def resolve_document_path(data_root: Path, document_path: str) -> Path:
    path = Path(document_path)

    if path.is_absolute():
        return path

    # Excel paths are expected to look like:
    # data/raw/policy/BIS_001.pdf
    if path.parts and path.parts[0] == "data":
        path = Path(*path.parts[1:])

    return data_root / path


def build_document(record: dict[str, str], file_path: Path, content: str) -> KnowledgeDocument:
    country_code = record.get("country_code", "")
    jurisdiction = record.get("jurisdiction", "")

    country_region = " / ".join(
        part for part in [country_code, jurisdiction] if part
    )

    header_fields = [
        ("Evidence ID", record.get("evidence_id", "")),
        ("Title", record.get("source_title", "")),
        ("Publisher", record.get("publisher", "")),
        ("Country/Jurisdiction", country_region),
        ("Publication date", record.get("publication_date", "")),
        ("Effective date", record.get("effective_date", "")),
        ("Status", record.get("status", "")),
        ("Topic", record.get("topic", "")),
        ("Applicability scope", record.get("applicability_scope", "")),
        ("Related industries", record.get("related_industries", "")),
        ("Claim", record.get("claim", "")),
        ("Key finding", record.get("key_finding", "")),
        ("Content summary", record.get("content_summary", "")),
        ("Applicability notes", record.get("applicability_notes", "")),
        ("Uncertainty", record.get("uncertainty", "")),
    ]

    # Drop empty fields and avoid storing the same text twice (the workbook
    # repeats key_finding in content_summary for most rows).
    metadata_lines: list[str] = []
    seen_values: set[str] = set()
    for label, value in header_fields:
        value = (value or "").strip()
        if not value:
            continue
        dedupe_key = value.lower()
        if label not in {"Evidence ID", "Title"} and dedupe_key in seen_values:
            continue
        seen_values.add(dedupe_key)
        metadata_lines.append(f"{label}: {value}")

    metadata_context = "\n".join(metadata_lines)

    combined_content = f"{metadata_context}\n\n{content}".strip()

    return KnowledgeDocument(
        source_id=record["evidence_id"],
        title=record.get("source_title", ""),
        source_type=record.get("source_type", ""),
        publisher=record.get("publisher", ""),
        country_region=country_region,
        publication_date=record.get("publication_date", ""),
        url=record.get("source_url", ""),
        authority_level=record.get("authority_level", ""),
        topic=record.get("topic", ""),
        document_path=str(file_path),
        content=combined_content,
        is_mock=False,
    )


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Ingest the curated EV/battery evidence dataset into the local RAG knowledge base."
    )

    parser.add_argument(
        "--data-root",
        type=Path,
        required=True,
        help="Path to the data directory containing metadata/ and raw/.",
    )

    parser.add_argument(
        "--excel",
        default="metadata/evidence.xlsx",
        help="Evidence metadata workbook path relative to data-root.",
    )
    parser.add_argument(
        "--evidence-id",
        action="append",
        default=[],
        help="Only ingest the specified evidence_id. Repeat for multiple IDs.",
    )

    args = parser.parse_args()

    data_root = args.data_root.expanduser().resolve()
    excel_path = data_root / args.excel

    if not excel_path.exists():
        print(f"ERROR: metadata file not found: {excel_path}")
        return 1

    settings = Settings.from_env()
    repository = KnowledgeRepository(settings.database_path)

    records = load_metadata(excel_path)
    if args.evidence_id:
        requested_ids = set(args.evidence_id)
        records = [
            record
            for record in records
            if record.get("evidence_id") in requested_ids
        ]

    print(f"Loaded {len(records)} metadata records.")
    print(f"Knowledge DB: {settings.database_path}")

    success = 0
    skipped = 0
    failed = 0

    for record in records:
        evidence_id = record.get("evidence_id", "UNKNOWN")

        # Do not ingest evidence explicitly marked as no longer current.
        status = record.get("status", "").upper()

        if status in {"EXPIRED", "REPEALED"}:
            print(f"[SKIP] {evidence_id}: status={status}")
            skipped += 1
            continue

        document_path = record.get("document_path", "")

        if not document_path:
            print(f"[SKIP] {evidence_id}: missing document_path")
            skipped += 1
            continue

        file_path = resolve_document_path(data_root, document_path)

        if not file_path.exists():
            print(f"[ERROR] {evidence_id}: file not found: {file_path}")
            failed += 1
            continue

        try:
            content = extract_content(file_path)

            if not content.strip():
                print(f"[SKIP] {evidence_id}: no extractable text")
                skipped += 1
                continue

            document = build_document(record, file_path, content)
            chunk_count = repository.ingest_document(document)

            print(
                f"[OK] {evidence_id}: "
                f"{file_path.name} -> {chunk_count} chunks"
            )

            success += 1

        except Exception as exc:
            print(f"[ERROR] {evidence_id}: {exc}")
            failed += 1

    print()
    print("=== INGESTION COMPLETE ===")
    print(f"Success: {success}")
    print(f"Skipped: {skipped}")
    print(f"Failed:  {failed}")
    print(f"Repository stats: {repository.stats()}")

    return 0 if failed == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
