from pathlib import Path
from openpyxl import load_workbook
import argparse


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def main():
    parser = argparse.ArgumentParser(
        description="Fix document_path values by matching evidence_id to files under data/raw."
    )
    parser.add_argument("--data-root", type=Path, required=True)
    parser.add_argument("--excel", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    data_root = args.data_root.expanduser().resolve()
    excel_path = data_root / args.excel
    output_path = data_root / args.output
    raw_root = data_root / "raw"

    if not excel_path.exists():
        print(f"ERROR: Excel not found: {excel_path}")
        return 1

    # 不使用 read_only，因为我们要修改单元格
    workbook = load_workbook(excel_path)

    sheet = (
        workbook["evidence"]
        if "evidence" in workbook.sheetnames
        else workbook.active
    )

    # 找到表头
    headers = {}

    for cell in sheet[1]:
        header = clean(cell.value)
        if header:
            headers[header] = cell.column

    if "evidence_id" not in headers:
        print("ERROR: evidence_id column not found.")
        return 1

    if "document_path" not in headers:
        print("ERROR: document_path column not found.")
        return 1

    evidence_col = headers["evidence_id"]
    path_col = headers["document_path"]

    # 建立 evidence_id -> 实际文件 的索引
    file_index = {}

    for path in raw_root.rglob("*"):
        if not path.is_file():
            continue

        if path.name == ".gitkeep":
            continue

        evidence_id = path.stem
        file_index.setdefault(evidence_id, []).append(path)

    fixed = 0
    already_correct = 0
    missing = 0
    ambiguous = 0

    for row in range(2, sheet.max_row + 1):
        evidence_id = clean(sheet.cell(row=row, column=evidence_col).value)
        document_path = clean(sheet.cell(row=row, column=path_col).value)

        if not evidence_id:
            continue

        # 先检查 Excel 原路径是否已经正确
        expected = None

        if document_path:
            relative_path = Path(document_path)

            # Excel 中路径以 data/ 开头时去掉 data/
            if relative_path.parts and relative_path.parts[0] == "data":
                relative_path = Path(*relative_path.parts[1:])

            expected = data_root / relative_path

        if expected and expected.exists():
            print(f"[OK] {evidence_id}")
            already_correct += 1
            continue

        candidates = file_index.get(evidence_id, [])

        # 只有唯一匹配时才自动修改，避免改错
        if len(candidates) == 1:
            actual_file = candidates[0]
            relative = actual_file.relative_to(data_root)

            new_document_path = f"data/{relative.as_posix()}"

            print(f"[FIX] {evidence_id}")
            print(f"  Old: {document_path or '(empty)'}")
            print(f"  New: {new_document_path}")

            sheet.cell(
                row=row,
                column=path_col
            ).value = new_document_path

            fixed += 1

        elif len(candidates) == 0:
            print(f"[MISSING] {evidence_id}")
            print(f"  Kept unchanged: {document_path or '(empty)'}")
            missing += 1

        else:
            print(f"[AMBIGUOUS] {evidence_id}")
            print("  Multiple matching files found:")

            for candidate in candidates:
                print(f"    {candidate}")

            print("  Not modified.")
            ambiguous += 1

    # 创建输出目录（如果需要）
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # 保存成新文件，不覆盖 v0.3
    workbook.save(output_path)

    print()
    print("=== PATH FIX COMPLETE ===")
    print(f"Already correct: {already_correct}")
    print(f"Fixed:           {fixed}")
    print(f"Missing:         {missing}")
    print(f"Ambiguous:       {ambiguous}")
    print()
    print(f"New Excel saved to:")
    print(output_path)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())