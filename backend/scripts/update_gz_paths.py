from pathlib import Path
from openpyxl import load_workbook
import argparse


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def evidence_id_from_file(path: Path) -> str:
    """
    Correctly handle normal files and double extensions.

    FR_001.json.gz -> FR_001
    TRD_004.json.gz -> TRD_004
    BIS_001.pdf -> BIS_001
    """
    name = path.name

    if name.lower().endswith(".json.gz"):
        return name[:-8]  # remove ".json.gz"

    return path.stem


def main():
    parser = argparse.ArgumentParser(
        description="Update document_path values for .json.gz files."
    )

    parser.add_argument("--data-root", type=Path, required=True)
    parser.add_argument("--excel", required=True)
    parser.add_argument("--output", required=True)

    args = parser.parse_args()

    data_root = args.data_root.expanduser().resolve()
    excel_path = data_root / args.excel
    output_path = data_root / args.output
    raw_root = data_root / "raw"

    if not excel_path.exists():
        print(f"ERROR: Excel not found: {excel_path}")
        return 1

    if not raw_root.exists():
        print(f"ERROR: raw directory not found: {raw_root}")
        return 1

    workbook = load_workbook(excel_path)

    sheet = (
        workbook["evidence"]
        if "evidence" in workbook.sheetnames
        else workbook.active
    )

    # Find columns
    headers = {}

    for cell in sheet[1]:
        header = clean(cell.value)
        if header:
            headers[header] = cell.column

    if "evidence_id" not in headers:
        print("ERROR: evidence_id column not found.")
        return 1

    if "document_path" not in headers:
        print("ERROR: document_path column not found.")
        return 1

    evidence_col = headers["evidence_id"]
    path_col = headers["document_path"]

    # Build an index ONLY for actual .json.gz files
    gz_index = {}

    for path in raw_root.rglob("*.json.gz"):
        if not path.is_file():
            continue

        evidence_id = evidence_id_from_file(path)
        gz_index.setdefault(evidence_id, []).append(path)

    updated = 0
    unchanged = 0
    ambiguous = 0

    for row in range(2, sheet.max_row + 1):
        evidence_id = clean(
            sheet.cell(row=row, column=evidence_col).value
        )

        old_path = clean(
            sheet.cell(row=row, column=path_col).value
        )

        if not evidence_id:
            continue

        candidates = gz_index.get(evidence_id, [])

        if len(candidates) == 1:
            actual_file = candidates[0]
            relative = actual_file.relative_to(data_root)

            new_path = f"data/{relative.as_posix()}"

            if old_path != new_path:
                print(f"[UPDATE] {evidence_id}")
                print(f"  Old: {old_path or '(empty)'}")
                print(f"  New: {new_path}")

                sheet.cell(
                    row=row,
                    column=path_col
                ).value = new_path

                updated += 1
            else:
                print(f"[OK] {evidence_id}: already .json.gz")
                unchanged += 1

        elif len(candidates) > 1:
            print(f"[AMBIGUOUS] {evidence_id}")

            for candidate in candidates:
                print(f"  {candidate}")

            ambiguous += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    print()
    print("=== GZ PATH UPDATE COMPLETE ===")
    print(f"Updated:   {updated}")
    print(f"Unchanged: {unchanged}")
    print(f"Ambiguous: {ambiguous}")
    print(f"Saved to:  {output_path}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())