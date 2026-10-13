"""Turn the Task 5 case database into ingestable corpus files.

Task 5 shipped ``data/processed/case_database_v0.1.xlsx`` (17 cases).  The RAG
ingest pipeline expects a raw document per evidence id plus a metadata row, so
this script materialises both:

  data/raw/case/<case_id>.md          structured case narrative
  data/metadata/evidence_cases_v0.1.xlsx   metadata workbook (Task 1 schema)

Metadata is derived deterministically from the case record; the derivation
rules are written into the workbook so they can be audited.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# case_id -> (home country, host country) actually involved in the decision.
GEOPOLITICS: dict[str, tuple[str, str]] = {
    "RELOC_001": ("CN", "TH"),
    "RELOC_002": ("CN", "TH"),
    "RELOC_003": ("CN", "VN"),
    "RELOC_004": ("CN", "MY"),
    "RELOC_005": ("CN", "VN"),
    "RELOC_006": ("CN", "TH"),
    "RELOC_007": ("CN", "ID"),
    "RELOC_008": ("CN", "ID"),
    "PLUS1_001": ("CN", "HU"),
    "PLUS1_002": ("CN", "ID"),
    "PLUS1_003": ("CN", "ID"),
    "PLUS1_004": ("CN", "MY/VN/US"),
    "PLUS1_005": ("KR", "VN"),
    "RETURN_001": ("CN", "MY"),
    "RETURN_002": ("TW/CN", "IN"),
    "RETURN_003": ("CN", "VN/TH/IN"),
    "RETURN_004": ("CN", "DE"),
}

# Task 5 `industry` -> Task 1 style `related_industries` tag.
INDUSTRY_TAG: dict[str, str] = {
    "Automotive (EV)": "整车/汽车",
    "EV Battery": "动力电池",
    "EV Battery / Energy Storage": "动力电池/储能",
    "EV Battery / Consumer Battery": "动力电池/消费电池",
    "Battery Materials": "电池材料",
    "Solar Manufacturing": "光伏",
    "Electronics / Semiconductor": "电子/半导体",
    "Electronics Manufacturing": "电子制造",
    "General Manufacturing / Supply Chain": "通用制造",
}

CASE_TYPE_CN: dict[str, str] = {
    "China → Southeast Asia relocation": "中国 → 东南亚转移",
    "China + 1 strategy": "China + 1 战略",
    "Reconsideration / Return to China": "重新评估 / 回流中国",
}

TRIGGER_CN: dict[str, str] = {
    "tariff_pressure": "关税压力",
    "geopolitical_uncertainty": "地缘政治不确定性",
    "cost_change": "成本变化",
    "supplier_issue": "供应商问题",
    "market_access": "市场准入",
    "regulation_change": "监管变化",
}

DERIVATION_RULES = [
    ("evidence_id", "Task 5 `case_id`", "直接沿用，作为 DB source_id"),
    ("source_type", "-", "统一为「案例」（Task 1 source_type 词表）"),
    ("publisher", "Task 5 `source`", "直接沿用发布/转载渠道"),
    ("country_code", "任务书 geopolitical mapping", "公司母国/案例涉及东道国，多国用 / 分隔"),
    ("topic", "-", "统一为「供应链布局」，与 Task 1 的 CASE_001–CASE_008 口径一致"),
    ("authority_level", "Task 5 `source`", "公司公告 / 投资者互动平台 / 媒体转述的公司官方表态 → S；媒体（含路透、财新、澎湃、财经网）→ B+"),
    ("verification_status", "Task 5 `source` + `verification_status`", "一手来源 → 已核实；媒体转述 → 已核实（二手来源）；原文待核 → 待核实原文"),
    ("related_industries", "Task 5 `industry`", "按 INDUSTRY_TAG 映射"),
    ("publication_date", "Task 5 `date`", "直接沿用"),
    ("source_url", "Task 5 `url`", "直接沿用"),
]


def clean(value) -> str:
    return "" if value is None else str(value).strip()


# A source is treated as a company-official filing when the channel itself is
# the company's (公告 / 投资者互动平台), or when a media outlet is explicitly
# relaying the company's own official statement (「（转××官方）」).  Media
# relaying another *media* outlet stays B+.
PRIMARY_MARKERS = ("公告", "投资者互动", "公司官方", "官方发布")


def is_primary_source(source: str) -> bool:
    if any(marker in source for marker in PRIMARY_MARKERS):
        return True
    return "（转" in source and "官方" in source


def grade_of(verification: str) -> str:
    text = clean(verification)
    if text.startswith("高"):
        return "高"
    if text.startswith("中高"):
        return "中高"
    return "中"


def derive_status(source: str, verification: str) -> str:
    if is_primary_source(source):
        return "已核实"
    if grade_of(verification) == "中":
        return "待核实原文"
    return "已核实（二手来源）"


def derive_authority(source: str) -> str:
    # Task 4 evidence taxonomy: "Official company filing" sits at S; media at B+.
    return "S" if is_primary_source(source) else "B+"


def case_title(case: dict) -> str:
    company = clean(case["company_name"])
    case_type = CASE_TYPE_CN.get(clean(case["case_type"]), clean(case["case_type"]))
    decision = clean(case["decision_type"])
    return f"{company}：{case_type}（{decision}）"


def build_narrative(case: dict) -> str:
    case_id = clean(case["case_id"])
    home, host = GEOPOLITICS.get(case_id, ("CN", ""))
    industry = clean(case["industry"])
    trigger = clean(case["decision_trigger"])

    sections: list[str] = []

    sections.append(
        "\n".join(
            [
                f"# {case_title(case)}",
                "",
                f"案例编号：{case_id}",
                f"企业：{clean(case['company_name'])}",
                f"行业：{industry}",
                f"案例类型：{CASE_TYPE_CN.get(clean(case['case_type']), clean(case['case_type']))}",
                f"母国/东道国：{home} → {host}",
                f"决策类型：{clean(case['decision_type'])}",
                f"发布时间：{clean(case['date'])}",
            ]
        )
    )

    sections.append(
        "\n".join(
            [
                "## 一、决策前状态（Initial Situation）",
                f"生产布局：{clean(case['initial_production_layout'])}",
                f"主要市场：{clean(case['initial_main_markets'])}",
                f"供应链结构：{clean(case['initial_supply_chain'])}",
                f"面临的问题：{clean(case['initial_problem'])}",
            ]
        )
    )

    sections.append(
        "\n".join(
            [
                "## 二、决策触发（Decision Trigger）",
                f"触发类型：{trigger}（{TRIGGER_CN.get(trigger, trigger)}）",
                f"触发细节：{clean(case['trigger_detail'])}",
            ]
        )
    )

    sections.append(
        "\n".join(
            [
                "## 三、实际决策（Actual Decision）",
                f"决策分类：{clean(case['decision_type'])}",
                f"决策内容：{clean(case['actual_decision'])}",
            ]
        )
    )

    sections.append(
        "\n".join(
            [
                "## 四、相关风险（Related Risk）",
                f"风险类别：{clean(case['related_risk'])}",
            ]
        )
    )

    sections.append(
        "\n".join(
            [
                "## 五、后续结果（Outcome）",
                f"成本影响：{clean(case['outcome_cost_impact'])}",
                f"生产表现：{clean(case['outcome_production_performance'])}",
                f"市场准入：{clean(case['outcome_market_access'])}",
                f"供应链影响：{clean(case['outcome_supply_chain_impact'])}",
                f"结果判定：{clean(case['outcome_verdict'])}",
            ]
        )
    )

    sections.append(
        "\n".join(
            [
                "## 六、来源与核验（Source & Verification）",
                f"来源：{clean(case['source'])}",
                f"原始链接：{clean(case['url'])}",
                f"核验状态：{clean(case['verification_status'])}",
                f"关联证据编号：{clean(case['related_evidence_ids'])}",
            ]
        )
    )

    if clean(case.get("notes")):
        sections.append("\n".join(["## 七、备注", clean(case["notes"])]))

    return "\n\n".join(sections).strip() + "\n"


def build_metadata_row(case: dict) -> dict:
    case_id = clean(case["case_id"])
    home, host = GEOPOLITICS.get(case_id, ("CN", ""))
    country_code = "/".join(p for p in (home, host) if p)
    source = clean(case["source"])
    verification = clean(case["verification_status"])
    industry = clean(case["industry"])

    claim = clean(case["actual_decision"])
    return {
        "evidence_id": case_id,
        "source_title": case_title(case),
        "source_type": "案例",
        "publisher": source,
        "country_code": country_code,
        "jurisdiction": "",
        "publication_date": clean(case["date"]),
        "effective_date": "",
        "status": "CURRENT",
        "source_url": clean(case["url"]),
        "authority_level": derive_authority(source),
        "topic": "供应链布局",
        "applicability_scope": "GENERAL",
        "related_industries": INDUSTRY_TAG.get(industry, industry),
        "verification_status": derive_status(source, verification),
        "retrieved_at": "",
        "document_path": f"data/raw/case/{case_id}.md",
        "claim": claim[:600],
        "key_finding": clean(case["outcome_verdict"]),
        "content_summary": (
            f"{clean(case['company_name'])}：{clean(case['initial_problem'])} "
            f"触发因素 {clean(case['decision_trigger'])}；"
            f"最终 {clean(case['decision_type'])}。{clean(case['outcome_verdict'])}"
        ),
        "applicability_notes": (
            f"适用于 {industry} 企业的海外产能布局决策模拟；"
            f"案例类型 {clean(case['case_type'])}。"
        ),
        "uncertainty": verification,
        "ev_battery_relevance": "高" if "Battery" in industry or "EV" in industry else "中",
        "screening_action": "KEEP",
        "screening_note": "Task 5 案例库，用于 company intelligence 与 scenario simulation 评测。",
    }


COLUMNS = [
    "evidence_id", "source_title", "source_type", "publisher", "country_code",
    "jurisdiction", "publication_date", "effective_date", "status", "source_url",
    "authority_level", "topic", "applicability_scope", "related_industries",
    "verification_status", "retrieved_at", "document_path", "claim",
    "key_finding", "content_summary", "applicability_notes", "uncertainty",
    "ev_battery_relevance", "screening_action", "screening_note",
]


def style_header(ws, ncol, color="1F3864"):
    fill = PatternFill("solid", fgColor=color)
    for c in ws[1][:ncol]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = "A2"


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--case-db",
        type=Path,
        default=Path("/Users/zhengyuanyuan/Desktop/project-data/data/processed/case_database_v0.1.xlsx"),
    )
    parser.add_argument(
        "--data-root",
        type=Path,
        default=Path("/Users/zhengyuanyuan/Desktop/project-data/data"),
    )
    args = parser.parse_args()

    workbook = load_workbook(args.case_db, read_only=True, data_only=True)
    sheet = workbook["Case Database"]
    rows = sheet.iter_rows(values_only=True)
    header = [clean(v) for v in next(rows)]
    cases = [dict(zip(header, row)) for row in rows if row and row[0]]
    workbook.close()

    raw_dir = args.data_root / "raw" / "case"
    raw_dir.mkdir(parents=True, exist_ok=True)

    for case in cases:
        case_id = clean(case["case_id"])
        path = raw_dir / f"{case_id}.md"
        path.write_text(build_narrative(case), encoding="utf-8")

    case_ids = [clean(c["case_id"]) for c in cases]
    unknown = [c for c in case_ids if c not in GEOPOLITICS]
    if unknown:
        raise SystemExit(f"missing geography mapping for: {unknown}")

    meta_dir = args.data_root / "metadata"
    meta_dir.mkdir(parents=True, exist_ok=True)
    out_path = meta_dir / "evidence_cases_v0.1.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "sources"
    ws.append(COLUMNS)
    for case in cases:
        row = build_metadata_row(case)
        ws.append([row.get(col, "") for col in COLUMNS])
    style_header(ws, len(COLUMNS))
    for i, col in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(46, len(col) + 12))

    rules = wb.create_sheet("derivation_rules")
    rules.append(["field", "Task 5 source", "rule"])
    for row in DERIVATION_RULES:
        rules.append(list(row))
    style_header(rules, 3, "7A3E9D")
    for i, w in enumerate([20, 34, 78], 1):
        rules.column_dimensions[get_column_letter(i)].width = w

    audit = wb.create_sheet("case_index")
    audit.append(["case_id", "company", "industry", "case_type", "decision_type",
                  "trigger", "country_code", "authority_level", "verification_status",
                  "source", "url"])
    for case in cases:
        row = build_metadata_row(case)
        audit.append([row["evidence_id"], clean(case["company_name"]), clean(case["industry"]),
                      clean(case["case_type"]), clean(case["decision_type"]),
                      clean(case["decision_trigger"]), row["country_code"],
                      row["authority_level"], row["verification_status"],
                      row["publisher"], row["source_url"]])
    style_header(audit, 11, "2E5E4E")

    wb.save(out_path)
    print(f"wrote {len(cases)} raw files to {raw_dir}")
    print(f"wrote metadata workbook to {out_path}")
    print(json.dumps({c: build_metadata_row(x)["authority_level"] for c, x in zip(case_ids, cases)},
                     ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
