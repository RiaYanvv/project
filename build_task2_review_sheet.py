"""Build a compact human-review sheet for Task 2 manual scoring.

Only rows where the AI-assisted human score differs from the automatic score are
included. Rows are grouped by how much judgement the product owner needs to
spend: A = straightforward, B1 = quick confirm, B2 = please decide.
"""
from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SRC = Path("Task2_manual_scoring_v3.xlsx")
OUT = Path("Task2_人工复核清单.xlsx")
DELIVERY = Path("Task2_交付物") / OUT.name

GROUP = {
    "Q002/CASE_002": "A-可直接确认",
    "Q002/CASE_003": "A-可直接确认",
    "Q010/WH_001": "A-可直接确认",
    "Q016/CASE_003": "A-可直接确认",
    "Q016/WH_002": "A-可直接确认",
    "Q028/DOC_001": "A-可直接确认",
    "Q030/WH_001": "A-可直接确认",
    "Q034/IRS_001": "A-可直接确认",
    "Q040/IRS_001": "A-可直接确认",
    "Q041/EU_003": "A-可直接确认",
    "Q043/EU_003": "A-可直接确认",
    "Q002/MOFCOM_004": "B1-快速确认",
    "Q012/WTO_002": "B1-快速确认",
    "Q018/MOFCOM_003": "B1-快速确认",
    "Q025/MOFCOM_003": "B1-快速确认",
    "Q025/MOFCOM_004": "B1-快速确认",
    "Q029/WTO_001": "B1-快速确认",
    "Q029/WTO_002": "B1-快速确认",
    "Q031/EU_003": "B1-快速确认",
    "Q031/UST_001": "B1-快速确认",
    "Q033/DHS_001": "B1-快速确认",
    "Q038/IRS_001": "B1-快速确认",
    "Q005/CASE_001": "B2-请你拍板",
    "Q014/NIO_001": "B2-请你拍板",
    "Q015/NIO_001": "B2-请你拍板",
    "Q023/USMCA_001": "B2-请你拍板",
    "Q031/SNE_001": "B2-请你拍板",
    "Q034/SNE_001": "B2-请你拍板",
    "Q035/IRS_001": "B2-请你拍板",
    "Q035/WTO_004": "B2-请你拍板",
    "Q038/EU_001": "B2-请你拍板",
    "Q038/SNE_001": "B2-请你拍板",
    "Q041/CSIS_001": "B2-请你拍板",
    "Q042/TRD_001": "B2-请你拍板",
    "Q043/EU_001": "B2-请你拍板",
}

HEADERS = [
    "序号", "复核分组", "题号", "排名", "Evidence ID", "证据标题", "发布机构",
    "权威级别", "Auto 相关性", "我的判定", "变化", "我的理由（节选）",
    "测试问题", "Gold（参考答案）", "用户判定 (0/1/2)", "用户备注",
]
WIDTHS = [6, 16, 8, 6, 14, 46, 30, 10, 11, 10, 9, 62, 55, 32, 14, 26]


def main() -> None:
    wb = load_workbook(SRC, data_only=True)
    ws = wb["Item Scoring"]
    h = {str(c.value): c.column for c in ws[1]}

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        note = str(r[h["Notes"] - 1] or "")
        if "人工判定与Auto" not in note:
            continue
        qid = str(r[h["Question ID"] - 1]).strip()
        eid = str(r[h["Retrieved Evidence ID"] - 1]).strip()
        key = f"{qid}/{eid}"
        if key not in GROUP:
            raise KeyError(f"ungrouped review row: {key}")
        auto = r[h["Auto Relevance (0-2)"] - 1]
        human = r[h["Relevance (0/1/2)"] - 1]
        rows.append({
            "group": GROUP[key],
            "qid": qid,
            "rank": r[h["Rank"] - 1],
            "eid": eid,
            "title": r[h["Title"] - 1],
            "publisher": r[h["Publisher"] - 1],
            "auth": r[h["Authority (S/A+/A/B+)"] - 1],
            "auto": auto,
            "human": human,
            "delta": f"{auto} → {human}",
            "note": note.split("｜ 人工判定与Auto")[0].strip(),
            "question": r[h["Question"] - 1],
            "gold": r[h["Expected Evidence IDs"] - 1],
        })

    order = {"A-可直接确认": 0, "B1-快速确认": 1, "B2-请你拍板": 2}
    rows.sort(key=lambda x: (order[x["group"]], x["qid"], x["rank"]))

    out = Workbook()
    sh = out.active
    sh.title = "复核清单"

    title = sh.cell(row=1, column=1, value="Task 2 人工评分复核清单（仅列出与 Auto 不一致的 35 行）")
    title.font = Font(bold=True, size=14)
    sh.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    sh.row_dimensions[1].height = 24

    sub = sh.cell(row=2, column=1, value=(
        "你只需处理 B2（黄色）几行；A（绿色）和 B1（浅蓝）如无异议可留空，视为采纳我的判定。"
        "如需改分，在“用户判定”列填 0 / 1 / 2，并在备注写一句理由。"
    ))
    sub.font = Font(size=10, color="555555")
    sh.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADERS))

    header_row = 4
    for c, (name, width) in enumerate(zip(HEADERS, WIDTHS), start=1):
        cell = sh.cell(row=header_row, column=c, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sh.column_dimensions[get_column_letter(c)].width = width
    sh.row_dimensions[header_row].height = 32

    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fills = {
        "A-可直接确认": PatternFill("solid", fgColor="E2EFDA"),
        "B1-快速确认": PatternFill("solid", fgColor="DDEBF7"),
        "B2-请你拍板": PatternFill("solid", fgColor="FFF2CC"),
    }

    for i, item in enumerate(rows, start=1):
        r = header_row + i
        values = [
            i, item["group"], item["qid"], item["rank"], item["eid"], item["title"],
            item["publisher"], item["auth"], item["auto"], item["human"], item["delta"],
            item["note"], item["question"], item["gold"], None, None,
        ]
        for c, v in enumerate(values, start=1):
            cell = sh.cell(row=r, column=c, value=v)
            cell.border = border
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=c in (6, 7, 12, 13, 16),
                horizontal="center" if c in (1, 2, 3, 4, 5, 8, 9, 10, 11, 15) else "left",
            )
            if c == 2:
                cell.fill = fills[item["group"]]
        sh.row_dimensions[r].height = 62

    last = header_row + len(rows)
    dv = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True,
                        showErrorMessage=True, errorTitle="只能填 0/1/2",
                        error="0=无关，1=部分相关，2=直接回答")
    sh.add_data_validation(dv)
    dv.add(f"O{header_row + 1}:O{last}")

    sh.conditional_formatting.add(
        f"O{header_row + 1}:O{last}",
        FormulaRule(formula=[f'AND($O{header_row + 1}<>"",$O{header_row + 1}<>$J{header_row + 1})'],
                    fill=PatternFill("solid", fgColor="FCE4D6")),
    )
    sh.freeze_panes = f"C{header_row + 1}"
    sh.auto_filter.ref = f"A{header_row}:P{last}"

    # ---- instruction sheet ----
    info = out.create_sheet("使用说明")
    info.column_dimensions["A"].width = 100
    lines = [
        ("Task 2 人工评分复核：你需要做什么", True),
        ("", False),
        ("1. 本表只收录“我的判定”与“Auto”不一致的 35 行（总表 215 行里的争议项）。", False),
        ("2. 绿色 A（11 行）：证据与问题直接对应，建议你扫一眼即可，留空=采纳。", False),
        ("3. 浅蓝 B1（11 行）：属于“部分相关”的边界判断，通常留空即可。", False),
        ("4. 黄色 B2（13 行）：跨主体/跨辖区，最值得你花时间拍板。", False),
        ("   —— 你只需在这 13 行的“用户判定”列填 0/1/2；不同意我时写一句理由。", False),
        ("   —— 留空 = 采纳“我的判定”。", False),
        ("5. 填完把文件发我，我会写回评分总表、重算 30 题指标，并同步交付目录。", False),
        ("", False),
        ("评分口径：", True),
        ("2 = 证据直接回答问题的核心信息需求；1 = 只覆盖其中一个维度/提供背景；0 = 跨主体、跨辖区或主题无关。", False),
        ("", False),
        ("当前 30 题口径（Q001–Q030）：Relevance 1.633 / Authority 1.847 / Hit@5 100%，均已达标。", False),
        ("复核不会轻易改变达标结论：即使 B2 全部按最保守方向调整，30 题 Relevance 仍高于 1.5。", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        cell = info.cell(row=i, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            cell.font = Font(bold=True, size=12 if i == 1 else 11)

    out.save(OUT)
    DELIVERY.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUT, DELIVERY)

    from collections import Counter
    counts = Counter(x["group"] for x in rows)
    print(f"saved {OUT} and {DELIVERY}")
    print(f"rows={len(rows)}  " + "  ".join(f"{k}={v}" for k, v in sorted(counts.items())))


if __name__ == "__main__":
    main()
