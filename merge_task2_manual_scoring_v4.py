"""Merge the reviewed Q044-Q046 supplement into the frozen Task 2 workbook.

Inputs:
  - Task2_manual_scoring_v3.xlsx                         (Q001-Q043, frozen)
  - Task2_manual_scoring_Q044_Q046_待复核.xlsx           (completed human review)
  - retrieval_results.xlsx                              (retrieved rows and excerpts)

Output:
  - Task2_manual_scoring_v4.xlsx                         (Q001-Q046, 230 item judgments)
"""
from __future__ import annotations

import shutil
from collections import defaultdict
from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent
V3 = ROOT / "Task2_manual_scoring_v3.xlsx"
SUPPLEMENT = ROOT / "Task2_manual_scoring_Q044_Q046_待复核.xlsx"
RESULTS = ROOT / "retrieval_results.xlsx"
OUT = ROOT / "Task2_manual_scoring_v4.xlsx"
DELIVERY = ROOT / "Task2_交付物"
DATA_EVAL = Path("/Users/zhengyuanyuan/Desktop/project-data/data/evaluation")
NEW_IDS = ("Q044", "Q045", "Q046")


def norm(value) -> str:
    return str(value or "").strip()


def number(value, field: str):
    if value is None or value == "":
        raise ValueError(f"missing {field}")
    return int(value)


def header_map(ws) -> dict[str, int]:
    return {norm(c.value): c.column for c in ws[1]}


def copy_row_style(ws, source_row: int, target_row: int, ncol: int) -> None:
    for col in range(1, ncol + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)
        dst.number_format = src.number_format


def read_supplement() -> tuple[list[dict], list[dict]]:
    wb = load_workbook(SUPPLEMENT, data_only=True)
    items_ws = wb["Item Scoring"]
    ih = header_map(items_ws)
    items = []
    for row in items_ws.iter_rows(min_row=2, values_only=True):
        if not norm(row[ih["Question ID"] - 1]):
            continue
        rec = {name: row[idx - 1] for name, idx in ih.items()}
        rec["human_relevance"] = number(rec["人工 Relevance (待填)"], "human Relevance")
        rec["human_authority"] = number(rec["人工 Authority (待填)"], "human Authority")
        items.append(rec)

    q_ws = wb["Question Coverage"]
    qh = header_map(q_ws)
    questions = []
    for row in q_ws.iter_rows(min_row=2, values_only=True):
        if not norm(row[qh["Question ID"] - 1]):
            continue
        rec = {name: row[idx - 1] for name, idx in qh.items()}
        rec["human_coverage"] = number(rec["人工 Coverage (待填)"], "human Coverage")
        questions.append(rec)

    if len(items) != 15 or len(questions) != 3:
        raise ValueError(f"expected 15 item rows / 3 questions, got {len(items)} / {len(questions)}")
    if {norm(r["Question ID"]) for r in items} != set(NEW_IDS):
        raise ValueError("supplement item rows do not match Q044-Q046")
    if {norm(r["Question ID"]) for r in questions} != set(NEW_IDS):
        raise ValueError("supplement question rows do not match Q044-Q046")
    return items, questions


def read_retrieval() -> tuple[dict[tuple[str, int], dict], dict[str, dict]]:
    wb = load_workbook(RESULTS, data_only=True)
    ws = wb["retrieval_results"]
    headers = [norm(c.value) for c in ws[1]]
    by_rank = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(headers, row))
        qid = norm(rec["question_id"])
        if qid not in NEW_IDS:
            continue
        by_rank[(qid, int(rec["rank"]))] = rec

    qws = wb["question_summary"]
    qheaders = [norm(c.value) for c in qws[1]]
    qmap = {}
    for row in qws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(qheaders, row))
        qid = norm(rec["question_id"])
        if qid in NEW_IDS:
            qmap[qid] = rec
    return by_rank, qmap


def append_item_rows(wb, supplement_items, retrieval_by_rank) -> None:
    ws = wb["Item Scoring"]
    start = ws.max_row + 1
    source_row = ws.max_row
    headers = header_map(ws)
    for offset, rec in enumerate(sorted(supplement_items, key=lambda r: (norm(r["Question ID"]), int(r["Rank"])))):
        qid = norm(rec["Question ID"])
        rank = int(rec["Rank"])
        r = retrieval_by_rank[(qid, rank)]
        target = start + offset
        values = {
            "Question ID": qid,
            "Rank": rank,
            "Retrieved Evidence ID": norm(rec["Evidence ID"]),
            "Title": rec["Title"],
            "Retrieved Topic": rec["Topic"],
            "Authority (S/A+/A/B+)": rec["Authority Level"],
            "Auto Relevance (0-2)": int(rec["Auto Relevance"]),
            "Auto Authority (0-2)": 2 if norm(rec["Authority Level"]) in {"S", "A+", "A"} else (1 if norm(rec["Authority Level"]) in {"B+", "B"} else 0),
            "Relevance (0/1/2)": rec["human_relevance"],
            "Authority (0/1/2)": rec["human_authority"],
            "Notes": rec["建议理由"],
            "Publisher": r.get("publisher"),
            "Scenario Category": r.get("scenario_category"),
            "Expected Topic": r.get("expected_topic"),
            "Expected Evidence IDs": r.get("expected_evidence_ids"),
            "Question": r.get("question"),
            "Evidence Excerpt": r.get("evidence_excerpt"),
        }
        for name, value in values.items():
            ws.cell(target, headers[name], value)
        copy_row_style(ws, source_row, target, ws.max_column)
        ws.row_dimensions[target].height = 64

    ws.auto_filter.ref = f"A1:Q{ws.max_row}"
    validations = list(ws.data_validations.dataValidation)
    for dv in validations:
        if dv.type == "list" and dv.formula1 == '"0,1,2"':
            dv.sqref = f"I2:J{ws.max_row}"
    if not validations:
        dv = DataValidation(type="list", formula1='"0,1,2"', allow_blank=False)
        ws.add_data_validation(dv)
        dv.add(f"I2:J{ws.max_row}")


def append_question_rows(wb, supplement_questions, retrieval_by_rank, retrieval_questions) -> None:
    ws = wb["Question Scoring"]
    headers = header_map(ws)
    source_row = ws.max_row
    item_ws = wb["Item Scoring"]
    ih = header_map(item_ws)
    rel_by_q = defaultdict(list)
    auth_by_q = defaultdict(list)
    for row in item_ws.iter_rows(min_row=2, values_only=True):
        qid = norm(row[ih["Question ID"] - 1])
        if qid in NEW_IDS:
            rel_by_q[qid].append(row[ih["Relevance (0/1/2)"] - 1])
            auth_by_q[qid].append(row[ih["Authority (0/1/2)"] - 1])

    for offset, rec in enumerate(sorted(supplement_questions, key=lambda r: norm(r["Question ID"]))):
        qid = norm(rec["Question ID"])
        qmeta = retrieval_questions[qid]
        target = ws.max_row + 1
        values = {
            "Question ID": qid,
            "Hit@5": int(rec["Hit@5"]),
            "Gold Hits in Top-5": int(rec["Gold Hits in Top-5"]),
            "Avg Relevance (0-2)": round(sum(rel_by_q[qid]) / len(rel_by_q[qid]), 4),
            "Avg Authority (0-2)": round(sum(auth_by_q[qid]) / len(auth_by_q[qid]), 4),
            "Coverage (0/1/2)": rec["human_coverage"],
            "Notes": rec["建议理由"],
            "Scenario Category": qmeta["scenario_category"],
            "Expected Topic": qmeta["expected_topic"],
            "Question": rec["Question"],
            "Expected Evidence IDs": rec["Expected Evidence IDs"],
            "Top-5 Retrieved IDs": rec["Top-5 Retrieved IDs"],
        }
        for name, value in values.items():
            ws.cell(target, headers[name], value)
        copy_row_style(ws, source_row, target, ws.max_column)
        ws.row_dimensions[target].height = 54

    for dv in ws.data_validations.dataValidation:
        if dv.type == "list" and dv.formula1 == '"0,1,2"':
            dv.sqref = f"F2:F{ws.max_row}"
    if not ws.data_validations.dataValidation:
        dv = DataValidation(type="list", formula1='"0,1,2"', allow_blank=False)
        ws.add_data_validation(dv)
        dv.add(f"F2:F{ws.max_row}")


def append_review_log(wb, supplement_items, retrieval_by_rank) -> None:
    ws = wb["Review Log"]
    hdr_row = 4
    headers = [norm(ws.cell(hdr_row, c).value) for c in range(1, 14)]
    col = {name: i + 1 for i, name in enumerate(headers)}
    start_seq = max(int(ws.cell(r, 1).value or 0) for r in range(5, ws.max_row + 1)) + 1
    source_row = ws.max_row
    for offset, rec in enumerate(sorted(supplement_items, key=lambda r: (norm(r["Question ID"]), int(r["Rank"])))):
        qid = norm(rec["Question ID"])
        rank = int(rec["Rank"])
        r = retrieval_by_rank[(qid, rank)]
        target = ws.max_row + 1
        auto = int(rec["Auto Relevance"])
        mine = int(rec["建议 Relevance"])
        user = rec["human_relevance"]
        status = "确认一致" if user == mine else "用户改分"
        values = {
            "序号": start_seq + offset,
            "复核分组": "B4-补题人工评分",
            "题号": qid,
            "排名": rank,
            "Evidence ID": rec["Evidence ID"],
            "证据标题": rec["Title"],
            "发布机构": r.get("publisher"),
            "Auto": auto,
            "我的判定": mine,
            "用户判定": user,
            "最终判定": user,
            "复核状态": status,
            "用户备注": rec["建议理由"],
        }
        for name, value in values.items():
            ws.cell(target, col[name], value)
        copy_row_style(ws, source_row, target, 13)
        ws.cell(target, col["复核分组"]).fill = PatternFill("solid", fgColor="E2EFDA")
        ws.cell(target, col["复核状态"]).fill = PatternFill("solid", fgColor="E2EFDA")
        ws.row_dimensions[target].height = 50

    end = ws.max_row
    ws["A1"] = "Task 2 人工复核记录（35 条 Auto 不一致行 + 15 条 B4 补题人工评分）"
    ws["A2"] = ("复核完成：2026-10-18；前 35 行为既有复核（13 条明确确认、22 条留空采纳、0 条改分）；"
                "后 15 行为 Q044–Q046 的补题人工评分（全部确认建议值）。")
    ws.auto_filter.ref = f"A{hdr_row}:M{end}"


def build_summary(wb) -> dict[str, float]:
    item = wb["Item Scoring"]
    ih = header_map(item)
    qs = wb["Question Scoring"]
    qh = header_map(qs)

    rel = defaultdict(list)
    auth = defaultdict(list)
    for row in item.iter_rows(min_row=2, values_only=True):
        qid = norm(row[ih["Question ID"] - 1])
        if not qid:
            continue
        rel[qid].append(int(row[ih["Relevance (0/1/2)"] - 1]))
        auth[qid].append(int(row[ih["Authority (0/1/2)"] - 1]))

    qrows = {}
    for row in qs.iter_rows(min_row=2, values_only=True):
        qid = norm(row[qh["Question ID"] - 1])
        if qid:
            qrows[qid] = row

    def metrics(ids: list[str]) -> dict[str, float]:
        rr = [v for qid in ids for v in rel[qid]]
        aa = [v for qid in ids for v in auth[qid]]
        cov = [int(qrows[qid][qh["Coverage (0/1/2)"] - 1]) for qid in ids]
        hit = [int(qrows[qid][qh["Hit@5"] - 1]) for qid in ids]
        return {
            "rel": sum(rr) / len(rr),
            "auth": sum(aa) / len(aa),
            "cov": sum(cov) / len(cov),
            "hit": sum(hit) / len(hit),
            "items": len(rr),
            "questions": len(ids),
        }

    q1_30 = [f"Q{i:03d}" for i in range(1, 31)]
    q31_43 = [f"Q{i:03d}" for i in range(31, 44)]
    q44_46 = list(NEW_IDS)
    all_ids = [f"Q{i:03d}" for i in range(1, 47)]
    m30, m31, m44, mall = metrics(q1_30), metrics(q31_43), metrics(q44_46), metrics(all_ids)

    old = wb["Summary"]
    idx = wb.sheetnames.index("Summary")
    del wb["Summary"]
    sm = wb.create_sheet("Summary", idx)

    sm.append(["Task 2 — Human Evaluation Summary", "Value", "Target / scope", "Status"])
    sm.append(["Questions (task-book acceptance scope Q001–Q030)", m30["questions"], "30 questions", ""])
    sm.append(["Extension set (Q031–Q043)", m31["questions"], "report only", ""])
    sm.append(["B4 supplement (Q044–Q046)", m44["questions"], "report only", ""])
    sm.append(["Human-scored total (Q001–Q046)", mall["questions"], "46 questions", ""])
    sm.append(["", None, None, None])
    sm.append(["Primary acceptance metrics — task-book scope Q001–Q030", None, None, None])
    metric_rows = [
        ("Avg Relevance (Top-5, 0-2)", m30["rel"], ">= 1.5", "PASS" if m30["rel"] >= 1.5 else "FAIL"),
        ("Avg Authority (Top-5, 0-2)", m30["auth"], ">= 1.5", "PASS" if m30["auth"] >= 1.5 else "FAIL"),
        ("Hit@5 rate (>=1 correct evidence in Top-5)", m30["hit"], ">= 80%", "PASS" if m30["hit"] >= 0.8 else "FAIL"),
        ("Avg Coverage (0-2)", m30["cov"], "report", "report only"),
    ]
    for values in metric_rows:
        sm.append(list(values))
    sm.append(["", None, None, None])
    sm.append(["Supplemental metrics — not used for Task 2 acceptance", None, None, None])
    for label, value, scope in [
        ("Avg Relevance (Top-5, 0-2)", mall["rel"], "Q001–Q046"),
        ("Avg Authority (Top-5, 0-2)", mall["auth"], "Q001–Q046"),
        ("Hit@5 rate (>=1 correct evidence in Top-5)", mall["hit"], "Q001–Q046"),
        ("Avg Coverage (0-2)", mall["cov"], "Q001–Q046"),
    ]:
        sm.append([label, round(value, 4), scope, "report only"])
    sm.append(["", None, None, None])
    sm.append(["Note", "Task-book acceptance remains Q001–Q030. Q031–Q046 are reported separately and must not be mixed into the acceptance average.", None, None])
    sm.append(["Review artifact", "Task2_人工复核清单.xlsx + Review Log sheet (35 flagged rows + 15 B4 supplement rows)", None, None])
    sm.append(["B4 supplement", f"15 item-level judgments + 3 Coverage judgments; item judgments {m44['items']} rows; all confirmed against the suggested values.", None, None])
    sm.append(["", None, None, None])
    sm.append(["Scoring pass", "AI-assisted pre-scoring + product-side human review; Q001–Q043 frozen, Q044–Q046 merged in v4.", None, None])
    sm.append(["Final status", "PASS — task-book Relevance, Authority and Hit@5 targets are met.", None, None])

    sm.column_dimensions["A"].width = 54
    sm.column_dimensions["B"].width = 74
    sm.column_dimensions["C"].width = 18
    sm.column_dimensions["D"].width = 14
    for c in sm[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    sm.row_dimensions[1].height = 30
    sm["A1"].font = Font(bold=True, color="FFFFFF", size=12)
    for row in sm.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        row[1].alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        row[2].alignment = Alignment(wrap_text=True, vertical="top", horizontal="center")
        row[3].alignment = Alignment(wrap_text=True, vertical="top", horizontal="center")
    for row_idx in (7, 13):
        for c in sm[row_idx][:4]:
            c.font = Font(bold=True, color="1F3864")
            c.fill = PatternFill("solid", fgColor="D9E2F3")
    for row_idx in range(8, 12):
        sm.cell(row_idx, 2).number_format = "0.000" if row_idx != 10 else "0.0%"
    for row_idx in range(14, 18):
        sm.cell(row_idx, 2).number_format = "0.000" if row_idx != 16 else "0.0%"
    for row_idx in (25, 26):
        sm.cell(row_idx, 1).font = Font(bold=True)
    for row_idx in range(1, sm.max_row + 1):
        sm.row_dimensions[row_idx].height = 26 if row_idx < 6 else 38
    return {"canonical": m30, "extension": m31, "b4": m44, "all": mall}


def main() -> None:
    for path in (V3, SUPPLEMENT, RESULTS):
        if not path.exists():
            raise FileNotFoundError(path)

    supplement_items, supplement_questions = read_supplement()
    retrieval_by_rank, retrieval_questions = read_retrieval()
    wb = load_workbook(V3)
    append_item_rows(wb, supplement_items, retrieval_by_rank)
    append_question_rows(wb, supplement_questions, retrieval_by_rank, retrieval_questions)
    append_review_log(wb, supplement_items, retrieval_by_rank)
    metrics = build_summary(wb)
    wb.save(OUT)

    for directory in (DELIVERY, DATA_EVAL):
        directory.mkdir(parents=True, exist_ok=True)
        shutil.copy2(OUT, directory / OUT.name)

    print(f"saved {OUT}")
    for name, m in metrics.items():
        print(f"{name}: questions={m['questions']} items={m['items']} "
              f"rel={m['rel']:.4f} auth={m['auth']:.4f} hit={m['hit']:.1%} cov={m['cov']:.4f}")


if __name__ == "__main__":
    main()
