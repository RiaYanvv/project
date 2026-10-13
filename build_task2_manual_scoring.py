"""Build the Task 2 human scoring workbook from retrieval_results.xlsx.

Output: Task2_manual_scoring_v3.xlsx
Sheets: Item Scoring (150 rows) / Question Scoring (30 rows) / Summary / Rubric
"""
from __future__ import annotations

import re
import sqlite3
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RESULTS = Path("retrieval_results.xlsx")
DB = Path("backend/test_data/app.db")
OUT = Path.home() / "Downloads" / "Task2_manual_scoring_v3.xlsx"

AUTHORITY_SCORE = {"S": 2, "A+": 2, "A": 2, "B+": 1, "B": 1}
TOPIC_MAP = {
    "supply_chain": {"供应链布局", "贸易流", "原产地与转运"},
    "supplier_dependency": {"供应链布局", "贸易流", "原产地与转运"},
    "export_control": {"出口管制", "出口管制/关税", "原产地与转运"},
    "trade_policy": {"原产地与转运", "反规避", "贸易流", "出口管制/关税", "供应链布局"},
    "investment_policy": {"对外投资", "宏观风险"},
    "market_access": {"原产地与转运", "反规避", "贸易流", "出口管制"},
    "regulation": {"出口管制", "出口管制/关税", "原产地与转运"},
}


def split_ids(value) -> list[str]:
    return [p.strip() for p in re.split(r"[;,，；]\s*", str(value or "")) if p.strip()]


def norm(eid: str) -> str:
    eid = str(eid or "").strip()
    if eid.startswith("EVD-"):
        eid = eid[4:]
    return re.sub(r"-C\d+$", "", eid)


def load_doc_meta() -> dict[str, dict]:
    con = sqlite3.connect(DB)
    meta = {}
    for sid, title, pub, auth, topic, date, url in con.execute(
        "select source_id,title,publisher,authority_level,topic,publication_date,url from knowledge_documents"
    ):
        meta[str(sid)] = {
            "title": (title or "").strip(),
            "publisher": (pub or "").strip(),
            "authority": (auth or "").strip(),
            "topic": (topic or "").strip(),
            "date": (date or "").strip(),
            "url": (url or "").strip(),
        }
    con.close()
    return meta


def trim(text, limit=600):
    text = re.sub(r"\s+", " ", str(text or "")).strip()
    return text if len(text) <= limit else text[: limit - 1] + "…"


def style_header(ws, ncol, color="1F3864"):
    fill = PatternFill("solid", fgColor=color)
    for c in ws[1][:ncol]:
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def main() -> None:
    wb_in = load_workbook(RESULTS, data_only=True)
    ws = wb_in["retrieval_results"]
    headers = [str(c.value) for c in ws[1]]
    rows = [dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    meta = load_doc_meta()

    wb = Workbook()

    # ---------------- Sheet 1: Item Scoring ----------------
    it = wb.active
    it.title = "Item Scoring"
    it.append([
        "Question ID", "Rank", "Retrieved Evidence ID", "Title", "Retrieved Topic",
        "Authority (S/A+/A/B+)", "Auto Relevance (0-2)", "Auto Authority (0-2)",
        "Relevance (0/1/2)", "Authority (0/1/2)", "Notes",
        "Publisher", "Scenario Category", "Expected Topic", "Expected Evidence IDs",
        "Question", "Evidence Excerpt",
    ])
    for r in rows:
        eid = norm(r["retrieved_evidence_id"])
        m = meta.get(eid, {})
        auth = str(r.get("authority_level") or m.get("authority") or "").strip()
        topic = str(r.get("retrieved_topic") or m.get("topic") or "").strip()
        exp = split_ids(r.get("expected_evidence_ids"))
        auto_rel = 2 if eid in exp else (1 if topic in TOPIC_MAP.get(str(r.get("expected_topic") or "").strip(), set()) else 0)
        it.append([
            r["question_id"], r["rank"], eid, m.get("title") or trim(r.get("title"), 120),
            topic, auth, auto_rel, AUTHORITY_SCORE.get(auth, ""),
            None, None, None,
            m.get("publisher") or r.get("publisher"), r["scenario_category"], r["expected_topic"],
            "; ".join(exp), r["question"], trim(r.get("evidence_excerpt")),
        ])
    nrow = it.max_row
    style_header(it, 17)
    widths = [11, 6, 18, 46, 15, 12, 10, 11, 12, 12, 46, 26, 20, 16, 24, 60, 70]
    for i, w in enumerate(widths, 1):
        it.column_dimensions[get_column_letter(i)].width = w
    thin = Side(style="thin", color="D9D9D9")
    human_fill = PatternFill("solid", fgColor="FFF2CC")
    for row in it.iter_rows(min_row=2, max_row=nrow, max_col=17):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for c in row[8:10]:
            c.fill = human_fill
    dv = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True)
    it.add_data_validation(dv)
    dv.add(f"I2:J{nrow}")
    it.auto_filter.ref = f"A1:Q{nrow}"
    it.freeze_panes = "D2"

    # ---------------- Sheet 2: Question Scoring ----------------
    qs = wb.create_sheet("Question Scoring")
    qs.append([
        "Question ID", "Hit@5", "Gold Hits in Top-5", "Avg Relevance (0-2)", "Avg Authority (0-2)",
        "Coverage (0/1/2)", "Notes",
        "Scenario Category", "Expected Topic", "Question", "Expected Evidence IDs", "Top-5 Retrieved IDs",
    ])
    order, seen = [], set()
    for r in rows:
        qid = r["question_id"]
        if qid not in seen:
            seen.add(qid)
            order.append(qid)
    by_q = {}
    for r in rows:
        by_q.setdefault(r["question_id"], []).append(r)
    for i, qid in enumerate(order):
        grp = by_q[qid]
        exp = split_ids(grp[0].get("expected_evidence_ids"))
        retr = [norm(r["retrieved_evidence_id"]) for r in grp]
        excel_row = i + 2
        hits = sum(1 for e in retr if e in exp)
        qs.append([
            qid, 1 if hits > 0 else 0, hits,
            f"=IFERROR(AVERAGEIF('Item Scoring'!$A$2:$A${nrow},$A{excel_row},'Item Scoring'!$I$2:$I${nrow}),\"\")",
            f"=IFERROR(AVERAGEIF('Item Scoring'!$A$2:$A${nrow},$A{excel_row},'Item Scoring'!$J$2:$J${nrow}),\"\")",
            None, None,
            grp[0]["scenario_category"], grp[0]["expected_topic"], grp[0]["question"],
            "; ".join(exp), "; ".join(retr),
        ])
    nq = qs.max_row
    style_header(qs, 12, color="2E5E4E")
    for i, w in enumerate([11, 8, 12, 12, 12, 12, 44, 20, 16, 60, 24, 40], 1):
        qs.column_dimensions[get_column_letter(i)].width = w
    for row in qs.iter_rows(min_row=2, max_row=nq, max_col=12):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for c in row[4:7]:
            c.fill = human_fill
    dv2 = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True)
    qs.add_data_validation(dv2)
    dv2.add(f"F2:F{nq}")
    qs.freeze_panes = "B2"

    # ---------------- Sheet 3: Summary ----------------
    sm = wb.create_sheet("Summary")
    sm.append(["Task 2 — Human Evaluation Summary", "Value", "Target", "Status"])
    sm.append(["Questions", 30, "30 questions", ""])
    sm.append(["", None, None, None])
    sm.append(["Primary metrics (auto-calculated from Question Scoring)", None, None, None])
    metrics = [
        ("Avg Relevance (Top-5, 0-2)", f"=AVERAGE('Question Scoring'!$D$2:$D${nq})", ">= 1.5", None),
        ("Avg Authority (Top-5, 0-2)", f"=AVERAGE('Question Scoring'!$E$2:$E${nq})", ">= 1.5", None),
        ("Hit@5 rate (>=1 correct evidence in Top-5)", f"=AVERAGE('Question Scoring'!$B$2:$B${nq})", ">= 80%", None),
        ("Avg Coverage (0-2)", f'=IFERROR(AVERAGE(\'Question Scoring\'!$F$2:$F${nq}),"")', "report", None),
    ]
    for name, formula, target, _ in metrics:
        sm.append([name, formula, target, None])
    first = sm.max_row - len(metrics) + 1
    for r in range(first, sm.max_row + 1):
        sm.cell(row=r, column=2).number_format = "0.000"
    sm.cell(row=first, column=4).value = f'=IF(B{first}="","",IF(B{first}>=1.5,"PASS","FAIL"))'
    sm.cell(row=first + 1, column=4).value = f'=IF(B{first + 1}="","",IF(B{first + 1}>=1.5,"PASS","FAIL"))'
    sm.cell(row=first + 2, column=4).value = f'=IF(B{first + 2}="","",IF(B{first + 2}>=0.8,"PASS","FAIL"))'
    sm.cell(row=first + 2, column=2).number_format = "0.0%"
    sm.cell(row=first + 3, column=4).value = "report only"
    sm.append(["", None, None, None])
    sm.append(["Note", "Fill the yellow cells first — Summary updates automatically.", None, None])
    sm.column_dimensions["A"].width = 46
    sm.column_dimensions["B"].width = 14
    sm.column_dimensions["C"].width = 16
    sm.column_dimensions["D"].width = 12
    for c in sm[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    sm["A1"].font = Font(bold=True, color="FFFFFF", size=12)

    # ---------------- Sheet 4: Rubric ----------------
    ru = wb.create_sheet("Rubric")
    ru.append(["Dimension", "Score", "Definition"])
    for row in [
        ("Relevance", 2, "Retrieved evidence directly answers / directly supports the question. (Usually: it is in Expected Evidence IDs.)"),
        ("Relevance", 1, "Partially relevant — same topic bucket, usable as supporting context but not a direct answer."),
        ("Relevance", 0, "Irrelevant to the question."),
        ("Authority", 2, "S / A+ / A source — government regulation, official company filing, WTO/UN/World Bank, IMF, CSIS, McKinsey, industry association."),
        ("Authority", 1, "B+ source — consulting reports, market research, brokerage / media summaries."),
        ("Authority", 0, "Low-quality or unverifiable source."),
        ("Coverage", 2, "Retrieval found enough evidence to answer the question end-to-end."),
        ("Coverage", 1, "Partial — some but not all facets covered."),
        ("Coverage", 0, "No usable supporting evidence retrieved."),
        ("", None, ""),
        ("How to score", None, "Score every retrieved item in 'Item Scoring' (yellow cells). Then give one Coverage score and Notes per question in 'Question Scoring'."),
        ("Auto columns", None, "Auto Relevance = 2 if the item is in Expected Evidence IDs, 1 if its topic is in the question's topic bucket, else 0. Auto Authority = 2 for S/A+/A, 1 for B+/B."),
    ]:
        ru.append(list(row))
    ru.column_dimensions["A"].width = 22
    ru.column_dimensions["B"].width = 8
    ru.column_dimensions["C"].width = 110
    for c in ru[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for row in ru.iter_rows(min_row=2, max_col=3):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"saved {OUT}")
    print(f"Item Scoring rows={nrow - 1}, Question Scoring rows={nq - 1}")


if __name__ == "__main__":
    main()
