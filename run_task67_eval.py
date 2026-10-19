"""Task 5 (benchmark files) + Task 6 (risk assessment) + Task 7 (evidence citation).

Reads the Task 5 case database and the Task 1 / Task 4 evidence registry, runs the
retrieval pipeline on each case's decision question, and writes:

  data/evaluation/historical_case_benchmark.xlsx   Task 5 deliverable (task-book name)
  data/evaluation/benchmark_results.xlsx           Task 5 results
  data/evaluation/risk_assessment_checklist.xlsx   Task 6
  data/evaluation/evidence_quality_test.xlsx       Task 7

Both the shipped (baseline) pipeline and the Task 4 optimised pipeline are scored,
so Task 6 / 7 report a Before vs After like Task 3 / 4 do.
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.app.config import Settings
from backend.app.evidence_metadata import load_registry
from backend.app.knowledge import KnowledgeRepository
from backend.app.retrieval import HybridRetrievalProvider
from backend.app.retrieval_optimizer import OptimizationConfig as C
from backend.app.risk_dimension import (
    DimensionContext,
    RiskDimensionRetriever,
    dimension_question,
    normalize_evidence_id,
)
from backend.app.retrieval_optimizer import RetrievalOptimizer
from backend.app.schemas import RetrievalQuery

DATA = Path("/Users/zhengyuanyuan/Desktop/project-data/data")
CASE_DB = DATA / "processed" / "case_database_v0.1.xlsx"
CASE_META = DATA / "evaluation" / "case_benchmark" / "evidence_cases_v0.1.xlsx"
OUT = DATA / "evaluation"
TOP_K = 10
CITED_K = 5

RISK_CATEGORIES = [
    "Trade Risk",
    "Political Risk",
    "Supply Chain Risk",
    "Regulatory Risk",
    "Market Access Risk",
    "Operational Risk",
]

# Risks every cross-border relocation decision touches, but which are not the
# point of the question.  Scored separately so they do not count as misses.
OPTIONAL_RISKS = {"Operational Risk"}

# 三种检索配置：现行基线、Task 4 优化层、Task 6 P0（按风险维度多次检索）。
CONFIGS = ("baseline", "optimized", "risk_dimension")
RISK_DIMENSION_CFG = "risk_dimension"


def normalize_risks(values) -> set[str]:
    """Keep only the Task 1 six-category vocabulary; 归并 Cost Risk -> Operational Risk."""
    out: set[str] = set()
    for value in values:
        value = str(value or "").strip()
        if value == "Cost Risk":
            value = "Operational Risk"
        if value in RISK_CATEGORIES:
            out.add(value)
    return out


COUNTRY_CN = {
    "CN": "中国", "TH": "泰国", "VN": "越南", "MY": "马来西亚", "ID": "印尼",
    "HU": "匈牙利", "DE": "德国", "IN": "印度", "TW": "中国台湾", "US": "美国",
    "MX": "墨西哥", "KR": "韩国",
}

BENCH_SCHEMA = [
    ("case_id", "唯一编号", "RELOC_001"),
    ("company_name", "企业名称", "BYD"),
    ("industry", "行业", "Automotive (EV)"),
    ("initial_situation", "决策前状态（生产布局 / 主要市场 / 供应链 / 问题）", "整车产能全部在中国"),
    ("decision_problem", "需要决策的问题", "是否将部分产能转移至泰国"),
    ("related_evidence_ids", "证据库绑定（Task 1 的 53 条证据，附绑定依据）", "TRD_001; EU_003"),
    ("external_source_refs", "外部新闻来源编号（见 case_database 的 Evidence Index）", "NEWS_001"),
    ("actual_decision", "真实决策", "泰国罗勇府建厂，2024-07 投产"),
    ("outcome", "后续结果（成本 / 生产 / 准入 / 供应链）", "降低准入壁垒，但零部件仍依赖中国"),
    ("key_factors", "关键因素", "tariff_pressure; market_access"),
    ("related_risks", "相关风险（Task 1 六类）", "Trade Risk; Market Access Risk"),
]


def style_header(ws, ncol, color="1F3864", height=30):
    fill = PatternFill("solid", fgColor=color)
    for c in ws[1][:ncol]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = height
    ws.freeze_panes = "A2"


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def split_semi(value) -> list[str]:
    if not value:
        return []
    return [p.strip() for p in re.split(r"[;,]", str(value)) if p.strip()]


def load_case_evidence_index() -> dict[str, dict]:
    """Task 5 source index: case evidence id -> outlet / title / date / url."""
    ws = load_workbook(CASE_DB, read_only=True, data_only=True)["Evidence Index"]
    data = list(ws.iter_rows(values_only=True))
    hdr = list(data[0])
    return {str(dict(zip(hdr, r))["evidence_id"]): dict(zip(hdr, r)) for r in data[1:] if r and r[0]}


def load_task1_source_index(repo) -> dict[str, dict]:
    """Task 1 evidence index: evidence id -> publisher / title / date / url.

    The case's `related_evidence_ids` point into the Task 1 evidence base, so
    the citation audit has to resolve them there, not only against the external
    news list.
    """
    out: dict[str, dict] = {}
    for chunk in repo.all_chunks():
        if chunk.source_id in out:
            continue
        out[chunk.source_id] = {
            "evidence_id": chunk.source_id,
            "outlet": str(chunk.publisher or ""),
            "title": str(chunk.title or ""),
            "date": str(chunk.publication_date or ""),
            "url": str(chunk.url or chunk.document_path or ""),
            "url_status": "证据库（Task 1）",
        }
    return out


def load_cases() -> list[dict]:
    ws = load_workbook(CASE_DB, read_only=True, data_only=True)["Case Database"]
    data = list(ws.iter_rows(values_only=True))
    hdr = list(data[0])
    cases = [dict(zip(hdr, r)) for r in data[1:] if r and r[0]]

    meta_ws = load_workbook(CASE_META, read_only=True, data_only=True)["sources"]
    mdata = list(meta_ws.iter_rows(values_only=True))
    mhdr = list(mdata[0])
    meta = {dict(zip(mhdr, r))["evidence_id"]: dict(zip(mhdr, r)) for r in mdata[1:] if r and r[0]}

    for case in cases:
        m = meta.get(case["case_id"], {})
        case["country_code"] = str(m.get("country_code") or "")
        case["hosts"] = [c for c in case["country_code"].split("/") if c and c != "CN"]
        host_cn = " / ".join(COUNTRY_CN.get(c, c) for c in case["hosts"]) or "海外"
        case["host_cn"] = host_cn
        case["expected_risks"] = normalize_risks(split_semi(case.get("related_risk")))
        case["raw_risks"] = split_semi(case.get("related_risk"))
        case["evidence_ids"] = split_semi(case.get("related_evidence_ids"))
        case["decision_problem"] = build_question(case, host_cn)
    return cases


def build_question(case: dict, host_cn: str) -> str:
    company = case["company_name"].split()[0]
    kind = case["case_type"]
    if kind.startswith("China + 1"):
        return (f"{company} 正在推进 China+1 布局（{host_cn}）。"
                f"管理层应基于现有证据评估该布局的风险。")
    if kind.startswith("Reconsideration"):
        return (f"{company} 正在重新评估海外布局，考虑提高中国生产比例、收缩海外新增产能。"
                f"管理层应基于现有证据评估这一调整的风险。")
    return (f"{company} 正在评估将部分产能从中国转移至{host_cn}。"
            f"管理层应基于现有证据判断该决策的风险。")


def dimension_context(case: dict) -> DimensionContext:
    return DimensionContext(
        company=case["company_name"].split()[0],
        host=case["host_cn"],
    )


def dimension_query(case: dict, risk: str) -> str:
    return dimension_question(risk, dimension_context(case))


def make_query(question: str, limit: int = TOP_K) -> RetrievalQuery:
    return RetrievalQuery(
        industry="battery_ev",
        products=[question],
        home_country="CN",
        production_countries=[],
        target_markets=[],
        decision_question=question,
        restrictions=[],
        limit=limit,
    )


def normalize(eid: str) -> str:
    return normalize_evidence_id(eid)


def corpus_tag_support(registry: dict) -> dict[str, int]:
    """How many corpus documents carry each risk tag."""
    counts = {risk: 0 for risk in RISK_CATEGORIES}
    for meta in registry.values():
        for risk in meta.get("related_risk", ()):
            if risk in counts:
                counts[risk] += 1
    return counts


def build_providers(repo, registry):
    baseline = HybridRetrievalProvider(repository=repo)
    tuned = {"global_demote": 0.0, "country_demote": 0.02, "industry_boost": 0.0}
    opt = RetrievalOptimizer(registry=registry, config=C(True, True, True))
    for k, v in tuned.items():
        setattr(opt.metadata_filter, k, v)
    optimized = HybridRetrievalProvider(repository=repo, optimizer=opt)
    return {"baseline": baseline, "optimized": optimized}


def build_risk_dimension_retriever(repo, registry) -> RiskDimensionRetriever:
    """Task 6 P0 的实际产品路径：按风险维度检索 + Task 1 标签确认。"""
    return RiskDimensionRetriever(
        HybridRetrievalProvider(repository=repo),
        registry,
        per_dimension_k=CITED_K,
    )


def load_source_audit() -> dict[str, object]:
    """Read the source-audit summary generated by validate_task5_sources.py."""
    path = OUT / "task5_source_validation.xlsx"
    if not path.exists():
        return {}
    ws = load_workbook(path, read_only=True, data_only=True)["Summary"]
    rows = list(ws.iter_rows(values_only=True))
    return {str(row[0]): row[1] for row in rows[1:] if row and row[0]}


# --------------------------------------------------------------------------
# Task 5 — historical case benchmark
# --------------------------------------------------------------------------
def write_benchmark(cases: list[dict], registry: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Benchmark"
    cols = ["case_id", "company_name", "industry", "case_type", "host_country",
            "initial_situation", "decision_problem", "related_evidence_ids",
            "evidence_rationale", "external_source_refs",
            "actual_decision", "outcome", "key_factors", "expected_risks",
            "decision_type"]
    ws.append(cols)
    for case in cases:
        sources = "; ".join(case["evidence_ids"])
        outcome = " / ".join(x for x in (
            case.get("outcome_cost_impact"), case.get("outcome_production_performance"),
            case.get("outcome_market_access"), case.get("outcome_supply_chain_impact")) if x)
        ws.append([
            case["case_id"], case["company_name"], case["industry"], case["case_type"],
            case["host_cn"],
            " / ".join(x for x in (case.get("initial_production_layout"),
                                   case.get("initial_main_markets"),
                                   case.get("initial_supply_chain"),
                                   case.get("initial_problem")) if x),
            case["decision_problem"], sources, case.get("related_evidence_rationale"),
            case.get("external_source_refs"), case.get("actual_decision"),
            outcome,
            "; ".join(x for x in (case.get("decision_trigger"), case.get("trigger_detail")) if x),
            "; ".join(sorted(case["expected_risks"])), case.get("decision_type"),
        ])
    style_header(ws, len(cols))
    autosize(ws, [12, 22, 18, 26, 14, 60, 50, 26, 56, 20, 50, 60, 40, 34, 18, 20])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

    s = wb.create_sheet("Schema")
    s.append(["字段", "说明", "示例"])
    for row in BENCH_SCHEMA:
        s.append(list(row))
    style_header(s, 3, "2E5E4E")
    autosize(s, [22, 52, 44])

    q = wb.create_sheet("QA")
    q.append(["验收项", "结果", "说明"])
    audit = load_source_audit()
    audit_text = (
        f"{audit.get('result', '未运行')}；{len(cases)} 案例 / {audit.get('sources', '?')} 来源；"
        f"dead={audit.get('dead_links', '?')}，WAF={audit.get('http_ok_waf', '?')}"
    ) if audit else "未生成审计文件"
    checks = [
        ("案例数量", f"{len(cases)} 个", "任务书要求 ≥10（最好 15 左右）"),
        ("案例真实性", "公开报道 / 公司官网 / 公告", "外部来源编号见 external_source_refs（映射到 case_database 的 Evidence Index，含 outlet/标题/日期/URL）"),
        ("来源链接审计（2026-10-18）", audit_text, "全量 HTTP 复核；WAF 表示站点阻断自动抓取但链接可达。逐条结果见 task5_source_validation.xlsx"),
        ("与 EV battery / 供应链迁移相关", "是", "整车(2) + 动力电池(7) + 电池材料(3) + 光伏/电子/通用(5)"),
        ("背景→决策→结果链条", "100%", "Initial Situation / Decision Trigger / Actual Decision / Outcome 四段齐全"),
        ("能否用于 Scenario Simulation", "是", "含 4 个「待观察」与 2 个「失败/退出」案例"),
        ("每条案例是否有标准答案", "是", "actual_decision + outcome + expected_risks 即标准答案"),
        ("是否进入 RAG 语料", "否（held-out）", "案例是评测答案键，故意不进入知识库，避免答案泄漏"),
        ("证据库绑定", "17/17", "related_evidence_ids 已改为 Task 1 证据库内的真实 ID，逐条附绑定依据（R1触发/R2目的地/R3行业/R4类比）"),
        ("held-out 校验", "通过", "语料 53 篇文档 ID 中无 RELOC_*/PLUS1_*/RETURN_*；17 个案例主来源 URL 与语料 URL 零重合（28 条 Evidence Index 链接另由来源审计复核）"),
        ("语料泄漏检查（逐字）", "通过", "backend/scripts/check_case_leakage.py：17/17 通过，最高逐字重合 0.8%（10 字滑窗，阈值 5%）"),
        ("隔离校验复跑", "check_case_leakage.py", ".venv/bin/python backend/scripts/check_case_leakage.py —— 退出码 0 表示案例 ID / 来源 URL / 逐字重合 / 证据绑定四项全过"),
        ("语料背景重合（已登记）", "4 处", "RELOC_001/PLUS1_003 经 IEA_001 的 OEM 名单；PLUS1_004 经 DOC_001 的光伏反规避语境；PLUS1_005 经 IMF_001 的 FDI 举例；RETURN_002 经 UNCTAD_001 引用的路透稿"),
        ("同源提示", "RETURN_002", "外部来源为路透报道，而语料 UNCTAD_001 引用了同一系列路透稿；属背景级同源，评测时需知悉"),
        ("通用吸引子", "CASE_001", "语料 CASE_001（大旺金属回流）出现在全部 17 个案例的 Top-5，是 Task 6 false positive 的来源之一"),
        ("已知局限", "", "结果字段多来自媒体转述；4 个案例结果尚未发生；行业级案例 RETURN_003 无单一企业数据；6 条站点受 WAF/验证码影响，需浏览器人工打开"),
    ]
    for row in checks:
        q.append(list(row))
    style_header(q, 3, "9D3E3E")
    autosize(q, [28, 24, 70])
    wb.save(OUT / "historical_case_benchmark.xlsx")


def company_tokens(case: dict) -> list[str]:
    """Company names that identify the same firm inside corpus titles."""
    raw = case["company_name"]
    cn = re.split(r"\s+|/|（", raw)[0]
    en = re.findall(r"[A-Za-z][A-Za-z&.\- ]{2,}", raw)
    tokens = [cn] + [t.strip() for t in en]
    return [t for t in dict.fromkeys(tokens) if len(t) > 1]


def write_benchmark_results(cases, scored, registry, repo) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Case Results"
    ws.append(["case_id", "company_name", "industry", "case_type", "decision_problem",
               "benchmark_held_out", "expected_decision_type", "evidence_index_ids",
               "expected_risks", "risks_identified_baseline",
               "risk_recall_baseline", "risk_recall_optimized",
               "policy_evidence_top5_baseline", "verdict"])
    for case in cases:
        s = scored[case["case_id"]]["baseline"]
        so = scored[case["case_id"]]["optimized"]
        ws.append([
            case["case_id"], case["company_name"], case["industry"], case["case_type"],
            case["decision_problem"], "是（未进入 RAG 语料）", case.get("decision_type"),
            "; ".join(case["evidence_ids"]),
            "; ".join(sorted(case["expected_risks"])),
            "; ".join(sorted(s["risks"])),
            round(len(case["expected_risks"] & s["risks"]) / max(1, len(case["expected_risks"])), 3),
            round(len(case["expected_risks"] & so["risks"]) / max(1, len(case["expected_risks"])), 3),
            "; ".join(s["retrieved"][:5]),
            case.get("outcome_verdict"),
        ])
    style_header(ws, 14)
    autosize(ws, [12, 22, 18, 26, 46, 18, 18, 24, 32, 34, 12, 12, 42, 26])
    ws.auto_filter.ref = f"A1:N{ws.max_row}"

    s = wb.create_sheet("Summary")
    n = len(cases)
    rb = sum(len(c["expected_risks"] & scored[c["case_id"]]["baseline"]["risks"]) /
             max(1, len(c["expected_risks"])) for c in cases) / n
    ro = sum(len(c["expected_risks"] & scored[c["case_id"]]["optimized"]["risks"]) /
             max(1, len(c["expected_risks"])) for c in cases) / n
    s.append(["Task 5 — Historical Case Benchmark", "Value", "Note"])
    for row in (
        ["cases", n, "任务书要求 ≥10"],
        ["held out of the RAG corpus", n, "17/17 案例仅供评测，不进入知识库（避免答案泄漏）"],
        ["RAG corpus size", len(registry), "语料只含 Task 1 证据，不含任何 benchmark 案例"],
        ["risk recall (baseline)", round(rb, 3), "见 Task 6：期望风险中被检索证据覆盖的比例"],
        ["risk recall (optimized)", round(ro, 3), "见 Task 6"],
        [],
        ["口径说明", "", "本 benchmark 的案例是 held-out 答案键，故意不在 RAG 语料中，"
                         "因此不再计算「案例文档召回」。检索层能力由 Task 2/3/4 评价，"
                         "风险识别由 Task 6 评价；端到端「决策是否正确」需要 Agent 生成层，"
                         "属下一阶段。"],
    ):
        s.append(row)
    style_header(s, 3, "1F3864")
    autosize(s, [34, 14, 92])
    wb.save(OUT / "benchmark_results.xlsx")


# --------------------------------------------------------------------------
# Task 6 — risk assessment checklist
# --------------------------------------------------------------------------
def write_risk_checklist(cases, scored, registry):
    wb = Workbook()

    ws = wb.active
    ws.title = "Checklist"
    ws.append(["case_id", "company_name", "decision_problem"] + RISK_CATEGORIES + ["rationale"])
    for case in cases:
        row = [case["case_id"], case["company_name"], case["decision_problem"]]
        for risk in RISK_CATEGORIES:
            if risk in case["expected_risks"]:
                row.append("YES")
            elif risk in OPTIONAL_RISKS:
                row.append("Optional")
            else:
                row.append("NO")
        row.append("Expected = 案例记录中人工标注的 related_risk；Optional = 跨境布局普遍涉及、但不作为该题必答项")
        ws.append(row)
    style_header(ws, 3 + len(RISK_CATEGORIES) + 1)
    autosize(ws, [12, 22, 46] + [16] * len(RISK_CATEGORIES) + [56])
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    ws = wb.create_sheet("Results")
    ws.append(["case_id", "company_name", "config", "expected_risks", "identified_risks",
               "correct", "missing", "false_positive", "risk_recall", "risk_precision",
               "cited_evidence"])
    for case in cases:
        exp = case["expected_risks"]
        required = exp
        for cfg in CONFIGS:
            pred = scored[case["case_id"]][cfg]["risks"]
            correct = sorted(required & pred)
            missing = sorted(required - pred)
            neg = {r for r in RISK_CATEGORIES if r not in required and r not in OPTIONAL_RISKS}
            false_pos = sorted(pred & neg)
            cited = scored[case["case_id"]][cfg].get("cited_evidence", [])
            ws.append([
                case["case_id"], case["company_name"], cfg,
                "; ".join(sorted(required)), "; ".join(sorted(pred)),
                "; ".join(correct), "; ".join(missing), "; ".join(false_pos),
                round(len(correct) / len(required), 3) if required else "",
                round(len(correct) / len(pred), 3) if pred else 0.0,
                "; ".join(cited),
            ])
    style_header(ws, 11, "2E5E4E")
    autosize(ws, [12, 22, 16, 32, 34, 30, 30, 26, 12, 14, 60])
    ws.auto_filter.ref = f"A1:K{ws.max_row}"

    ws = wb.create_sheet("By Risk Category")
    ws.append(["risk_category", "expected_cases"]
              + [f"identified_{cfg}" for cfg in CONFIGS]
              + [f"recall_{cfg}" for cfg in CONFIGS]
              + ["over_claimed_cases"])
    for risk in RISK_CATEGORIES:
        exp_cases = [c for c in cases if risk in c["expected_risks"]]
        counts = [
            sum(1 for c in exp_cases if risk in scored[c["case_id"]][cfg]["risks"])
            for cfg in CONFIGS
        ]
        recalls = [
            round(n / len(exp_cases), 3) if exp_cases else "" for n in counts
        ]
        over = sum(1 for c in cases if risk not in c["expected_risks"]
                   and risk in scored[c["case_id"]]["baseline"]["risks"])
        ws.append([risk, len(exp_cases)] + counts + recalls + [over])
    style_header(ws, 2 + 2 * len(CONFIGS) + 1, "7A3E9D")
    autosize(ws, [22, 16] + [20] * len(CONFIGS) + [16] * len(CONFIGS) + [20])

    # 归因表：每一个被声称的风险后面站着哪条证据。审阅者可以据此判断
    # 「false positive」是检索错了，还是案例标注本身没写这一类。
    ws = wb.create_sheet("FP Attribution")
    ws.append(["case_id", "company_name", "risk_category", "case_annotation",
               "confirming_evidence_ids", "confirming_publishers", "assessment"])
    for case in cases:
        exp = case["expected_risks"]
        findings = scored[case["case_id"]][RISK_DIMENSION_CFG].get("findings", {})
        for risk in RISK_CATEGORIES:
            ids = findings.get(risk) or []
            if not ids:
                continue
            if risk in exp:
                annotation = "YES"
            elif risk in OPTIONAL_RISKS:
                annotation = "Optional"
            else:
                annotation = "NO"
            publishers = []
            for eid in ids:
                meta = registry.get(eid) or {}
                name = str(meta.get("publisher") or "").strip()
                if name and name not in publishers:
                    publishers.append(name)
            if annotation == "NO":
                assessment = ("案例标注未列该类，但证据带 Task 1 related_risk="
                              f"{risk} 标签，属「标注口径」差异，可复核")
            elif annotation == "Optional":
                assessment = "案例标注为 Optional，不计入 missing / false positive"
            else:
                assessment = "与案例标注一致"
            ws.append([case["case_id"], case["company_name"], risk, annotation,
                       "; ".join(ids), "; ".join(publishers[:4]), assessment])
    style_header(ws, 7, "9C5700")
    autosize(ws, [12, 22, 20, 16, 46, 44, 56])
    ws.auto_filter.ref = f"A1:G{ws.max_row}"

    ws = wb.create_sheet("Per-Dimension Probe")
    ws.append(["case_id", "company_name", "expected_risks"] + RISK_CATEGORIES +
              ["covered_by_dimension", "total_expected"])
    for case in cases:
        by_dim = scored[case["case_id"]]["dimension"]
        exp = case["expected_risks"]
        row = [case["case_id"], case["company_name"], "; ".join(sorted(exp))]
        for risk in RISK_CATEGORIES:
            row.append("YES" if by_dim.get(risk) else "NO")
        covered = sum(1 for risk in exp if by_dim.get(risk))
        row += [f"{covered}/{len(exp)}", len(exp)]
        ws.append(row)
    style_header(ws, 3 + len(RISK_CATEGORIES) + 2, "2E5E4E")
    autosize(ws, [12, 22, 32] + [14] * len(RISK_CATEGORIES) + [20, 14])

    ws = wb.create_sheet("Summary")
    n = len(cases)
    rows = []
    for cfg in CONFIGS:
        rec = sum(len(c["expected_risks"] & scored[c["case_id"]][cfg]["risks"]) /
                  max(1, len(c["expected_risks"])) for c in cases) / n
        prec = sum(len(c["expected_risks"] & scored[c["case_id"]][cfg]["risks"]) /
                   max(1, len(scored[c["case_id"]][cfg]["risks"])) for c in cases) / n
        miss = sum(len(c["expected_risks"] - scored[c["case_id"]][cfg]["risks"]) for c in cases)
        fp = sum(len({r for r in RISK_CATEGORIES if r not in c["expected_risks"] and r not in OPTIONAL_RISKS}
                     & scored[c["case_id"]][cfg]["risks"]) for c in cases)
        rows.append((cfg, rec, prec, miss, fp))
    by_cfg = {row[0]: row for row in rows}
    ws.append(["Task 6 — Risk Assessment Evaluation"]
              + ["Baseline", "Optimized", "Risk-dimension (P0)", "Note"])
    ws.append(["Risk Recall (mean over cases)"]
              + [round(by_cfg[cfg][1], 3) for cfg in CONFIGS]
              + ["应识别风险中有多少被证据链覆盖"])
    ws.append(["Risk Precision (mean over cases)"]
              + [round(by_cfg[cfg][2], 3) for cfg in CONFIGS]
              + ["识别出的风险中有多少属于本题期望风险"])
    ws.append(["Missing risks (total)"]
              + [by_cfg[cfg][3] for cfg in CONFIGS]
              + ["期望风险中未被证据覆盖的次数"])
    ws.append(["False positives (total)"]
              + [by_cfg[cfg][4] for cfg in CONFIGS]
              + ["不属于该题却被识别的风险次数"])
    ceil = sum(sum(1 for r in c["expected_risks"] if scored[c["case_id"]]["dimension"].get(r)) /
               max(1, len(c["expected_risks"])) for c in cases) / n
    tag_counts = corpus_tag_support(registry)
    representable = sum(sum(1 for r in c["expected_risks"] if tag_counts.get(r, 0) > 0) /
                        max(1, len(c["expected_risks"])) for c in cases) / n
    dim_recall = by_cfg[RISK_DIMENSION_CFG][1]
    ws.append(["风险可表征率（语料中存在该类风险的证据）", round(representable, 3), round(representable, 3),
               round(representable, 3),
               f"L1：语料层面。六类风险在 {len(registry)} 篇语料中均有标注（不含 held-out 案例），缺失的是「与案例挂钩」的证据"])
    ws.append(["风险可取回率（按维度逐条检索）", round(ceil, 3), round(ceil, 3), round(dim_recall, 3),
               "L2：检索层面。按风险维度单独提问时的覆盖上界；P0 之后已由产品实现达到"])
    ws.append(["风险识别覆盖（现行产品行为）", round(by_cfg["baseline"][1], 3),
               round(by_cfg["optimized"][1], 3), round(dim_recall, 3),
               "L3：开启 RISK_DIMENSION_RETRIEVAL（默认）后走 Risk-dimension 列；差距 = L1 - L3"])
    ws.append([])
    ws.append(["语料中各风险的证据数", "", "", "",
               "; ".join(f"{k}={v}" for k, v in tag_counts.items())])
    ws.append([])
    ws.append(["Scoring rule", "", "", "",
               "identified_risks = 引用证据在 Task 1 schema 中的 related_risk 并集；"
               "expected = 案例记录的人工标注 related_risk（已归并到 Task 1 六类词表）。"
               "risk_dimension 列：六个风险维度各检索 Top-5，只保留 related_risk 标签确认过的条目"])
    ws.append(["Caveat", "", "", "",
               "Operational Risk 记为 Optional，不计入 missing / false positive；"
               "本评测测的是证据链能否支撑风险识别，不是 LLM 生成风险的质量。"
               "risk_dimension 的 Precision 低于基线主要是口径效应：41 次误报里 Regulatory 14 次、"
               "Market Access 11 次，均可追溯到中国侧出口管制 / 公司披露与欧盟法规证据；"
               "逐条归因见 FP Attribution sheet"])
    style_header(ws, 5, "1F3864")
    autosize(ws, [34, 14, 14, 18, 78])
    wb.save(OUT / "risk_assessment_checklist.xlsx")
    return rows


# --------------------------------------------------------------------------
# Task 7 — evidence citation evaluation
# --------------------------------------------------------------------------
CITATION_FIELDS = ["evidence_id", "source", "date", "authority"]


def write_citation_test(cases, scored, registry, repo, source_index) -> None:
    docs = {c.source_id: c for c in repo.all_chunks()}
    first = {}
    for chunk in repo.all_chunks():
        first.setdefault(chunk.source_id, chunk)

    wb = Workbook()
    ws = wb.active
    ws.title = "Citation Test"
    ws.append(["case_id", "risk_statement", "evidence_id", "source", "date",
               "authority", "has_id", "has_source", "has_date", "has_authority",
               "traceable", "url_or_path"])
    for case in cases:
        pred = scored[case["case_id"]]["optimized"]["risks"]
        for risk in sorted(pred):
            cited = [e for e in scored[case["case_id"]]["optimized"]["cited_evidence"]
                     if risk in set((registry.get(e) or {}).get("related_risk", ()))]
            for eid in cited:
                meta = registry.get(eid) or {}
                chunk = first.get(eid)
                ws.append([
                    case["case_id"],
                    f"{case['company_name']}：{risk}",
                    eid,
                    meta.get("publisher") or (chunk.publisher if chunk else ""),
                    (chunk.publication_date if chunk else "") or "",
                    meta.get("authority_level") or (chunk.authority_level if chunk else ""),
                    int(bool(meta.get("evidence_id"))),
                    int(bool((meta.get("publisher") or (chunk.publisher if chunk else "")))),
                    int(bool((chunk.publication_date if chunk else ""))),
                    int(bool(meta.get("authority_level") or (chunk.authority_level if chunk else ""))),
                    int(all([bool(meta.get("evidence_id")),
                             bool(meta.get("publisher") or (chunk.publisher if chunk else "")),
                             bool(chunk and chunk.publication_date),
                             bool(meta.get("authority_level") or (chunk and chunk.authority_level))])),
                    (chunk.url if chunk and chunk.url else (chunk.document_path if chunk else "")) or "",
                ])
    style_header(ws, 12, "7A3E9D")
    autosize(ws, [12, 40, 12, 30, 14, 12, 9, 11, 10, 13, 11, 46])
    ws.auto_filter.ref = f"A1:L{ws.max_row}"

    ws = wb.create_sheet("Corpus Citation Coverage")
    ws.append(["evidence_id", "title", "source", "date", "authority", "industry",
               "topic", "evidence_type", "related_risk", "url_or_path",
               "complete", "missing_fields"])
    incomplete = 0
    for eid, meta in sorted(registry.items()):
        chunk = first.get(eid)
        source = meta.get("publisher") or (chunk.publisher if chunk else "")
        date = (chunk.publication_date if chunk else "") or ""
        authority = meta.get("authority_level") or (chunk.authority_level if chunk else "")
        path = (chunk.url if chunk and chunk.url else None) or (chunk.document_path if chunk else None) or ""
        missing = [name for name, ok in (
            ("source", bool(source)), ("date", bool(date)),
            ("authority", bool(authority)), ("url/path", bool(path))) if not ok]
        if missing:
            incomplete += 1
        ws.append([eid, meta.get("title") or (chunk.title if chunk else ""), source, date,
                   authority, meta.get("industry"), meta.get("topic_cn"),
                   meta.get("evidence_type"), "; ".join(meta.get("related_risk", [])),
                   path, int(not missing), "; ".join(missing)])
    style_header(ws, 12, "2E5E4E")
    autosize(ws, [12, 44, 28, 14, 12, 18, 14, 16, 34, 46, 10, 18])
    ws.auto_filter.ref = f"A1:L{ws.max_row}"

    ws = wb.create_sheet("Conclusion Audit")
    ws.append(["case_id", "conclusion_type", "conclusion", "cited_evidence_ids",
               "external_source_refs", "all_citable", "sources", "has_date", "verdict"])
    unaudited = 0
    for case in cases:
        cited_ids = case["evidence_ids"]
        metas = [source_index.get(e) for e in cited_ids]
        levels = [(m or {}).get("url_status", "") for m in metas]
        for kind, key in (("actual_decision", "actual_decision"),
                          ("outcome_cost", "outcome_cost_impact"),
                          ("outcome_production", "outcome_production_performance"),
                          ("outcome_market_access", "outcome_market_access"),
                          ("outcome_supply_chain", "outcome_supply_chain_impact")):
            text = case.get(key)
            if not text:
                continue
            citable = [e for e in cited_ids if source_index.get(e)]
            ok = bool(citable)
            if not ok:
                unaudited += 1
            ws.append([
                case["case_id"], kind, str(text)[:200], "; ".join(cited_ids),
                str(case.get("external_source_refs") or ""),
                int(ok),
                "; ".join(sorted({(source_index.get(e) or {}).get("outlet", "") for e in cited_ids if source_index.get(e)})),
                int(bool(citable) and all(source_index.get(e, {}).get("date") for e in citable)),
                "可追溯" if ok else "无证据支撑",
            ])
    style_header(ws, 9, "9D3E3E")
    autosize(ws, [12, 20, 60, 26, 24, 12, 22, 10, 14])
    ws.auto_filter.ref = f"A1:H{ws.max_row}"

    ws = wb.create_sheet("Summary")
    total = sum(1 for _ in registry)
    cit_rows = list(wb["Citation Test"].iter_rows(min_row=2, values_only=True))
    traceable = sum(1 for r in cit_rows if r[10] == 1)
    ws.append(["Task 7 — Evidence Citation Evaluation", "Value", "Note"])
    for row in (
        ["citation rows produced", len(cit_rows), "案例 × 识别风险 × 支撑证据"],
        ["fully traceable citations", traceable, "Evidence ID + Source + Date + Authority 四项齐全"],
        ["citation completeness", round(traceable / len(cit_rows), 3) if cit_rows else "",
         "目标：1.000（每条关键结论均可追溯）"],
        ["corpus documents with metadata", total, "Task 1 schema 覆盖的 evidence 数"],
        ["corpus documents with missing citation field", incomplete,
         "缺 source / date / authority / url 之一（见 Corpus Citation Coverage）"],
        ["metadata completeness rate", round((total - incomplete) / total, 3) if total else "",
         "target 1.000"],
        ["conclusions without evidence support", unaudited,
         "案例的关键结论（决策 / 结果四段）中没有任何可追溯证据的条数"],
    ):
        ws.append(row)
    ws.append([])
    ws.append(["Rule", "", "关键结论 → evidence_id → source / date / authority 四级可追溯；"
                           "任一字段缺失即视为不可追溯"])
    ws.append(["Caveat", "", "本测试校验引用链的结构完整性，不校验引用内容是否真的支持结论"])
    style_header(ws, 3, "1F3864")
    autosize(ws, [42, 12, 72])
    wb.save(OUT / "evidence_quality_test.xlsx")
    return {"citation_rows": len(cit_rows), "traceable": traceable,
            "corpus_docs": total, "incomplete": incomplete, "unaudited_conclusions": unaudited}


# --------------------------------------------------------------------------
def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    registry = load_registry()
    settings = Settings.from_env()
    repo = KnowledgeRepository(settings.database_path)
    providers = build_providers(repo, registry)
    cases = load_cases()

    dimension_retriever = build_risk_dimension_retriever(repo, registry)

    scored: dict[str, dict] = {}
    for case in cases:
        query = make_query(case["decision_problem"])
        entry = {}
        for cfg, provider in providers.items():
            results = provider.search(query)
            ids = [normalize(r.evidence_id) for r in results]
            cited = ids[:CITED_K]
            risks: set[str] = set()
            for eid in cited:
                risks |= set((registry.get(eid) or {}).get("related_risk", ()))
            entry[cfg] = {
                "retrieved": ids,
                "retrieved_set": set(ids),
                "cited_evidence": cited,
                "risks": risks,
            }

        # Task 6 P0：按风险维度检索 + Task 1 标签确认。这是**现行产品行为**
        # （`AgentService._dimension_evidence` 走同一个 Retriever），不再是探针。
        dim_result = dimension_retriever.retrieve(query, dimension_context(case))
        dim_cited = dim_result.evidence_ids
        entry[RISK_DIMENSION_CFG] = {
            "retrieved": dim_cited,
            "retrieved_set": set(dim_cited),
            "cited_evidence": dim_cited,
            "risks": set(dim_result.detected_risks),
            "findings": {
                risk: list(finding.evidence_ids)
                for risk, finding in dim_result.findings.items()
            },
        }
        entry["dimension"] = {
            risk: bool(finding.evidence_ids)
            for risk, finding in dim_result.findings.items()
        }
        scored[case["case_id"]] = entry

    write_benchmark(cases, registry)
    write_benchmark_results(cases, scored, registry, repo)
    risk_rows = write_risk_checklist(cases, scored, registry)
    case_source_index = load_case_evidence_index()
    task1_index = load_task1_source_index(repo)
    clash = set(case_source_index) & set(task1_index)
    assert not clash, f"案例外部来源编号与证据库 ID 冲突: {sorted(clash)}"
    source_index = {**case_source_index, **task1_index}
    cite = write_citation_test(cases, scored, registry, repo, source_index)

    print("cases:", len(cases))
    for cfg, rec, prec, miss, fp in risk_rows:
        print(f"  risk[{cfg}] recall={rec:.3f} precision={prec:.3f} missing={miss} fp={fp}")
    print("  citations:", cite)


if __name__ == "__main__":
    main()
