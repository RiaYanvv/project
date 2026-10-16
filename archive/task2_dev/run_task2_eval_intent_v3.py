from __future__ import annotations

from dataclasses import replace

import csv
import json
import re
from pathlib import Path

from openpyxl import load_workbook

from backend.app.config import Settings
from backend.app.knowledge import KnowledgeRepository
from backend.app.retrieval import HybridRetrievalProvider
from backend.app.schemas import RetrievalQuery


BENCHMARK = Path("task2_retrieval_evaluation_v0.7.xlsx")
OUTPUT_JSON = Path("task2_retrieval_intent_v3.json")
OUTPUT_CSV = Path("task2_retrieval_intent_v3.csv")


def normalize_evidence_id(evidence_id: str) -> str:
    """
    Examples:
      EVD-CASE_005-C0619   -> CASE_005
      EVD-TRD_001-C0000    -> TRD_001
      EVD-MOFCOM_002-C0021 -> MOFCOM_002
    """
    value = evidence_id.strip()

    if value.startswith("EVD-"):
        value = value[4:]

    value = re.sub(r"-C\d+$", "", value)
    return value


def split_expected(value) -> list[str]:
    if value is None:
        return []

    text = str(value).strip()
    if not text:
        return []

    # Workbook currently uses semicolon-separated evidence IDs.
    parts = re.split(r"[;,，；]\s*", text)
    return [part.strip() for part in parts if part.strip()]


def load_questions(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)

    if "retrieval_questions" not in workbook.sheetnames:
        raise RuntimeError(
            "Workbook does not contain sheet 'retrieval_questions'."
        )

    sheet = workbook["retrieval_questions"]
    rows = sheet.iter_rows(values_only=True)

    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    questions = []

    for row in rows:
        record = dict(zip(headers, row))

        question_id = str(record.get("question_id") or "").strip()
        question = str(record.get("question") or "").strip()

        if not question_id or not question:
            continue

        questions.append(
            {
                "question_id": question_id,
                "question": question,
                "scenario_category": str(
                    record.get("scenario_category") or ""
                ).strip(),
                "expected_topic": str(
                    record.get("expected_topic") or ""
                ).strip(),
                "expected_evidence_ids": split_expected(
                    record.get("expected_evidence_ids")
                ),
                "difficulty": str(
                    record.get("difficulty") or ""
                ).strip(),
            }
        )

    return questions


def main() -> None:
    if not BENCHMARK.exists():
        raise FileNotFoundError(f"Benchmark not found: {BENCHMARK}")

    settings = replace(
        Settings.from_env(),
        database_path=Path("backend/test_data/app_task2_clean.db"),
    )
    repository = KnowledgeRepository(settings.database_path)

    # IMPORTANT:
    # Do not instantiate WebSearchTool.
    # This makes live_budget=0 and gives all five slots to local RAG.
    retrieval = HybridRetrievalProvider(
        repository=repository,
        fallback=None,
        web_search=None,
    )

    questions = load_questions(BENCHMARK)

    print(f"Loaded {len(questions)} benchmark questions.")
    print(f"Database: {settings.database_path}")
    print("Mode: local-only hybrid retrieval")
    print("Top-K: 5")
    print()

    all_results = []
    hit_count = 0

    for index, item in enumerate(questions, start=1):
        qid = item["question_id"]
        question = item["question"]
        gold = item["expected_evidence_ids"]

        # Neutral, consistent benchmark configuration.
        # Geography contained in the natural-language question is NOT
        # duplicated into structured retrieval hints.
        query = RetrievalQuery(
            industry="battery_ev",
            products=[question],
            home_country="CN",
            production_countries=[],
            target_markets=[],
            decision_question=question,
            restrictions=[],
            limit=5,
        )

        retrieved = retrieval.search(query)

        # Retrieval already limits each source_id to one result, but normalize
        # and deduplicate again at evidence/document level for evaluation.
        unique_results = []
        seen_ids = set()

        for result in retrieved:
            normalized_id = normalize_evidence_id(result.evidence_id)

            if normalized_id in seen_ids:
                continue

            seen_ids.add(normalized_id)

            unique_results.append(
                {
                    "rank": len(unique_results) + 1,
                    "raw_evidence_id": result.evidence_id,
                    "evidence_id": normalized_id,
                    "relevance_score_raw": result.relevance_score,
                    "authority_level": result.authority_level,
                    "publisher": result.publisher,
                    "topic": result.topic,
                    "title": result.title,
                    "content": result.content,
                }
            )

            if len(unique_results) >= 5:
                break

        retrieved_ids = [r["evidence_id"] for r in unique_results]
        matched_ids = [eid for eid in gold if eid in retrieved_ids]
        hit_at_5 = 1 if matched_ids else 0
        hit_count += hit_at_5

        record = {
            **item,
            "retrieved": unique_results,
            "retrieved_evidence_ids": retrieved_ids,
            "matched_evidence_ids": matched_ids,
            "hit_at_5": hit_at_5,
        }
        all_results.append(record)

        status = "HIT" if hit_at_5 else "MISS"

        print(
            f"[{index:02d}/{len(questions):02d}] "
            f"{qid}: {status} | "
            f"Gold={','.join(gold) or '-'} | "
            f"Top5={','.join(retrieved_ids) or '-'}"
        )

    total = len(all_results)
    hit_rate = hit_count / total if total else 0.0

    summary = {
        "benchmark_file": str(BENCHMARK),
        "retrieval_mode": "hybrid_local_only",
        "web_search": False,
        "top_k": 5,
        "industry": "battery_ev",
        "home_country": "CN",
        "production_countries": [],
        "target_markets": [],
        "question_count": total,
        "hit_count": hit_count,
        "miss_count": total - hit_count,
        "hit_at_5": hit_rate,
        "target_hit_at_5": 0.80,
        "target_passed": hit_rate > 0.80,
    }

    OUTPUT_JSON.write_text(
        json.dumps(
            {
                "summary": summary,
                "results": all_results,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    with OUTPUT_CSV.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as handle:
        fieldnames = [
            "question_id",
            "scenario_category",
            "difficulty",
            "expected_topic",
            "expected_evidence_ids",
            "rank",
            "retrieved_evidence_id",
            "raw_evidence_id",
            "retrieval_score",
            "authority_level",
            "publisher",
            "topic",
            "hit_at_5",
            "matched_evidence_ids",
        ]

        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()

        for item in all_results:
            for result in item["retrieved"]:
                writer.writerow(
                    {
                        "question_id": item["question_id"],
                        "scenario_category": item["scenario_category"],
                        "difficulty": item["difficulty"],
                        "expected_topic": item["expected_topic"],
                        "expected_evidence_ids": ";".join(
                            item["expected_evidence_ids"]
                        ),
                        "rank": result["rank"],
                        "retrieved_evidence_id": result["evidence_id"],
                        "raw_evidence_id": result["raw_evidence_id"],
                        "retrieval_score": result["relevance_score_raw"],
                        "authority_level": result["authority_level"],
                        "publisher": result["publisher"],
                        "topic": result["topic"],
                        "hit_at_5": item["hit_at_5"],
                        "matched_evidence_ids": ";".join(
                            item["matched_evidence_ids"]
                        ),
                    }
                )

    print()
    print("=" * 72)
    print("TASK 2 RETRIEVAL BASELINE")
    print("=" * 72)
    print(f"Questions : {total}")
    print(f"Hits      : {hit_count}")
    print(f"Misses    : {total - hit_count}")
    print(f"Hit@5     : {hit_rate:.1%}")
    print(f"Target    : >80%")
    print(
        "Status    : "
        + ("PASS" if hit_rate > 0.80 else "BELOW TARGET")
    )
    print()
    print(f"Raw JSON  : {OUTPUT_JSON}")
    print(f"CSV       : {OUTPUT_CSV}")


if __name__ == "__main__":
    main()
