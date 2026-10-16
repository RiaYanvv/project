from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from openpyxl import load_workbook

from backend.app.config import Settings
from backend.app.knowledge import KnowledgeRepository
from backend.app.retrieval import HybridRetrievalProvider
from backend.app.schemas import RetrievalQuery


BENCHMARK = Path("task2_retrieval_evaluation_v0.7.xlsx")
OUTPUT_CSV = Path("task2_baseline_results.csv")
OUTPUT_JSON = Path("task2_baseline_results.json")


def normalize_evidence_id(value: str) -> str:
    value = str(value or "").strip()

    if value.startswith("EVD-"):
        value = value[4:]

    return re.sub(r"-C\d+$", "", value)


def split_expected(value) -> list[str]:
    if value is None:
        return []

    text = str(value).strip()
    if not text:
        return []

    return [
        item.strip()
        for item in re.split(r"[;,，；]\s*", text)
        if item.strip()
    ]


def load_questions(path: Path) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)

    if "retrieval_questions" not in wb.sheetnames:
        raise RuntimeError(
            "Sheet 'retrieval_questions' not found. "
            f"Available sheets: {wb.sheetnames}"
        )

    ws = wb["retrieval_questions"]
    rows = ws.iter_rows(values_only=True)

    try:
        header_row = next(rows)
    except StopIteration:
        raise RuntimeError("retrieval_questions sheet is empty.")

    headers = [
        str(value).strip() if value is not None else ""
        for value in header_row
    ]

    required = {
        "question_id",
        "question",
        "expected_evidence_ids",
    }

    missing = required - set(headers)
    if missing:
        raise RuntimeError(
            f"Missing required columns: {sorted(missing)}\n"
            f"Found columns: {headers}"
        )

    questions = []

    for row in rows:
        record = dict(zip(headers, row))

        qid = str(record.get("question_id") or "").strip()
        question = str(record.get("question") or "").strip()

        if not qid or not question:
            continue

        questions.append(
            {
                "question_id": qid,
                "question": question,
                "expected_evidence_ids": split_expected(
                    record.get("expected_evidence_ids")
                ),
                "expected_topic": str(
                    record.get("expected_topic") or ""
                ).strip(),
                "scenario_category": str(
                    record.get("scenario_category") or ""
                ).strip(),
                "difficulty": str(
                    record.get("difficulty") or ""
                ).strip(),
            }
        )

    return questions


def main() -> None:
    if not BENCHMARK.exists():
        raise FileNotFoundError(BENCHMARK)

    settings = Settings.from_env()

    repository = KnowledgeRepository(settings.database_path)

    retrieval = HybridRetrievalProvider(
        repository=repository,
        fallback=None,
        web_search=None,
    )

    questions = load_questions(BENCHMARK)

    print("=" * 72)
    print("TASK 2 RETRIEVAL BASELINE")
    print("=" * 72)
    print(f"Benchmark: {BENCHMARK}")
    print(f"Database:  {settings.database_path}")
    print(f"Questions: {len(questions)}")
    print("Mode:      LOCAL ONLY")
    print("Web:       OFF")
    print("Top-K:     5")
    print("=" * 72)

    if len(questions) != 30:
        print(
            f"WARNING: expected 30 questions, "
            f"but loaded {len(questions)}."
        )

    output = []
    hit_count = 0

    for index, row in enumerate(questions, start=1):
        qid = row["question_id"]
        question = row["question"]
        gold = set(row["expected_evidence_ids"])

        query = RetrievalQuery(
            industry="battery_ev",
            products=[question],
            home_country="CN",
            production_countries=[],
            target_markets=[],
            decision_question=question,
            restrictions=[],
            limit=5,
        )

        results = retrieval.search(query)

        retrieved = []
        detailed_results = []

        for rank, item in enumerate(results, start=1):
            normalized = normalize_evidence_id(
                item.evidence_id
            )

            retrieved.append(normalized)

            detailed_results.append(
                {
                    "rank": rank,
                    "raw_evidence_id": item.evidence_id,
                    "evidence_id": normalized,
                    "relevance_score": item.relevance_score,
                    "topic": item.topic,
                    "authority_level": item.authority_level,
                    "publisher": item.publisher,
                    "title": item.title,
                }
            )

        matched = [
            evidence_id
            for evidence_id in retrieved
            if evidence_id in gold
        ]

        hit = 1 if matched else 0
        hit_count += hit

        result = {
            **row,
            "retrieved_ids": retrieved,
            "matched_ids": matched,
            "hit_at_5": hit,
            "results": detailed_results,
        }

        output.append(result)

        status = "HIT " if hit else "MISS"

        print(
            f"{index:02d}/30 {qid}: {status} | "
            f"gold={sorted(gold)} | "
            f"top5={retrieved} | "
            f"matched={matched}"
        )

    total = len(output)
    hit_rate = hit_count / total if total else 0.0

    print()
    print("=" * 72)
    print("FINAL BASELINE")
    print("=" * 72)
    print(f"Questions: {total}")
    print(f"Hits:      {hit_count}")
    print(f"Misses:    {total - hit_count}")
    print(f"Hit@5:     {hit_rate:.4f} ({hit_rate * 100:.2f}%)")
    print("=" * 72)

    OUTPUT_JSON.write_text(
        json.dumps(
            {
                "benchmark": str(BENCHMARK),
                "database": str(settings.database_path),
                "mode": "local-only",
                "top_k": 5,
                "question_count": total,
                "hits": hit_count,
                "misses": total - hit_count,
                "hit_at_5": hit_rate,
                "results": output,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    with OUTPUT_CSV.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as f:
        writer = csv.writer(f)

        writer.writerow(
            [
                "question_id",
                "question",
                "expected_evidence_ids",
                "retrieved_top5",
                "matched_ids",
                "hit_at_5",
                "expected_topic",
                "scenario_category",
                "difficulty",
            ]
        )

        for row in output:
            writer.writerow(
                [
                    row["question_id"],
                    row["question"],
                    "; ".join(
                        row["expected_evidence_ids"]
                    ),
                    "; ".join(row["retrieved_ids"]),
                    "; ".join(row["matched_ids"]),
                    row["hit_at_5"],
                    row["expected_topic"],
                    row["scenario_category"],
                    row["difficulty"],
                ]
            )

    print()
    print(f"Saved: {OUTPUT_CSV}")
    print(f"Saved: {OUTPUT_JSON}")


if __name__ == "__main__":
    main()
